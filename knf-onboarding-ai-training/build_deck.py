# -*- coding: utf-8 -*-
"""한전원자력연료 신입직원 입문교육 교재 (세로형) 생성기"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---------- THEME ----------
NAVY  = RGBColor(0x0F, 0x24, 0x44)
BLUE  = RGBColor(0x1B, 0x5E, 0xA8)
SKY   = RGBColor(0x3E, 0x8E, 0xDE)
TEAL  = RGBColor(0x12, 0xA1, 0x8E)
ORANGE= RGBColor(0xE8, 0x7A, 0x2B)
GOLD  = RGBColor(0xF2, 0xB1, 0x3C)
LIGHT = RGBColor(0xF3, 0xF6, 0xFB)
LINE  = RGBColor(0xD9, 0xE2, 0xEE)
INK   = RGBColor(0x1B, 0x27, 0x38)
GRAY  = RGBColor(0x5C, 0x6B, 0x80)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RED   = RGBColor(0xD2, 0x4B, 0x3E)
PALE_TEAL = RGBColor(0xE7, 0xF5, 0xF2)
PALE_ORG  = RGBColor(0xFC, 0xF0, 0xE4)
PALE_BLUE = RGBColor(0xE9, 0xF1, 0xFB)
PALE_RED  = RGBColor(0xFB, 0xEA, 0xE8)
CARDNAVY  = RGBColor(0x16, 0x33, 0x5C)
SUBSKY    = RGBColor(0x9F, 0xC7, 0xF0)
FONT = "맑은 고딕"

prs = Presentation()
prs.slide_width  = Inches(7.5)
prs.slide_height = Inches(10.833)
BLANK = prs.slide_layouts[6]
PW, PH = prs.slide_width, prs.slide_height
M = Inches(0.5)

def slide(): return prs.slides.add_slide(BLANK)

def _set_font(run, size, color, bold, font=FONT):
    run.font.size = Pt(size); run.font.color.rgb = color
    run.font.bold = bold; run.font.name = font
    rPr = run._r.get_or_add_rPr()
    for tag in ('a:ea','a:cs'):
        e = rPr.find(qn(tag))
        if e is None:
            e = rPr.makeelement(qn(tag), {}); rPr.append(e)
        e.set('typeface', font)

def rect(s, x, y, w, h, fill, line=None, line_w=0, shape=MSO_SHAPE.RECTANGLE, shadow=False, round_=None):
    sp = s.shapes.add_shape(shape, x, y, w, h)
    if fill is None: sp.fill.background()
    else: sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None: sp.line.fill.background()
    else: sp.line.color.rgb = line; sp.line.width = Pt(line_w)
    sp.shadow.inherit = False
    if shadow:
        el = sp._element.spPr
        ef = el.makeelement(qn('a:effectLst'), {})
        sh = el.makeelement(qn('a:outerShdw'), {'blurRad':'80000','dist':'38000','dir':'5400000','rotWithShape':'0'})
        clr = el.makeelement(qn('a:srgbClr'), {'val':'1B2738'})
        al = el.makeelement(qn('a:alpha'), {'val':'18000'})
        clr.append(al); sh.append(clr); ef.append(sh); el.append(ef)
    if round_ is not None and shape in (MSO_SHAPE.ROUNDED_RECTANGLE, MSO_SHAPE.ROUND_2_SAME_RECTANGLE):
        try: sp.adjustments[0]=round_
        except: pass
    return sp

def text(s, x, y, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space_after=2, line_spacing=1.0, wrap=True):
    tb = s.shapes.add_textbox(x, y, w, h); tf = tb.text_frame
    tf.word_wrap = wrap; tf.vertical_anchor = anchor
    tf.margin_left=0; tf.margin_right=0; tf.margin_top=0; tf.margin_bottom=0
    first=True
    for para in paras:
        p = tf.paragraphs[0] if first else tf.add_paragraph(); first=False
        p.alignment = align; p.space_after = Pt(space_after); p.space_before=Pt(0); p.line_spacing = line_spacing
        for r in para:
            run = p.add_run(); run.text = r[0]
            fnt = r[4] if len(r)>4 else FONT
            _set_font(run, r[1], r[2], r[3], fnt)
    return tb

def bg(s, color=WHITE): rect(s, 0, 0, PW, PH, color)

def header(s, tag, title, color, accent, page_no):
    rect(s, 0, 0, PW, Inches(0.15), accent)
    text(s, M, Inches(0.4), PW-2*M-Inches(0.5), Inches(0.3), [[(tag, 11, color, True)]])
    text(s, M, Inches(0.68), PW-2*M-Inches(0.4), Inches(0.85), [[(title, 21.5, NAVY, True)]], line_spacing=0.98)
    rect(s, M, Inches(1.46), Inches(0.85), Inches(0.055), accent)
    text(s, PW-Inches(1.05), Inches(0.44), Inches(0.55), Inches(0.3), [[(str(page_no), 12, GRAY, True)]], align=PP_ALIGN.RIGHT)

def chip(s, x, y, w, h, label, color, txtcolor=WHITE, size=10.5):
    rect(s, x, y, w, h, color, shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.5)
    text(s, x, y, w, h, [[(label, size, txtcolor, True)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

def def_box(s, y, label, body, accent, pale, h=Inches(1.12)):
    rect(s, M, y, PW-2*M, h, pale, shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.05)
    rect(s, M, y, Inches(0.13), h, accent, shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.5)
    text(s, Inches(0.8), y+Inches(0.15), Inches(6), Inches(0.35), [[(label, 12, accent, True)]])
    text(s, Inches(0.8), y+Inches(0.46), Inches(6.0), Inches(0.6), [[(body, 12, INK, False)]], line_spacing=1.18)
    return y+h

def usage_card(s, x, y, w, h, no, title, body, prompt, accent, pale):
    rect(s, x, y, w, h, WHITE, line=LINE, line_w=1, shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.035, shadow=True)
    chip(s, x+Inches(0.18), y+Inches(0.16), Inches(0.95), Inches(0.3), "활용 "+no, accent, size=10)
    text(s, x+Inches(1.25), y+Inches(0.14), w-Inches(1.4), Inches(0.34), [[(title, 12.5, NAVY, True)]], anchor=MSO_ANCHOR.MIDDLE)
    text(s, x+Inches(0.2), y+Inches(0.6), w-Inches(0.42), Inches(1.0), [[(body, 11, INK, False)]], line_spacing=1.16)
    py = y+h-Inches(0.94)
    rect(s, x+Inches(0.18), py, w-Inches(0.36), Inches(0.8), LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.07)
    text(s, x+Inches(0.32), py+Inches(0.08), w-Inches(0.6), Inches(0.25), [[("💬 프롬프트 예시", 9, accent, True)]])
    text(s, x+Inches(0.32), py+Inches(0.31), w-Inches(0.62), Inches(0.46), [[(prompt, 9.5, GRAY, False)]], line_spacing=1.05)

def steps(s, y, title, accent, rows, rh=Inches(0.78)):
    w = PW-2*M
    if title:
        text(s, M, y, w, Inches(0.4), [[(title, 13.5, NAVY, True)]]); y=y+Inches(0.44)
    for i,(t,d) in enumerate(rows):
        rect(s, M, y, w, rh, LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.06)
        rect(s, M+Inches(0.12), y+Inches(0.16), Inches(0.46), Inches(0.46), accent, shape=MSO_SHAPE.OVAL)
        text(s, M+Inches(0.12), y+Inches(0.16), Inches(0.46), Inches(0.46), [[(str(i+1), 14, WHITE, True)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(s, M+Inches(0.75), y+Inches(0.1), w-Inches(0.9), Inches(0.32), [[(t, 12, NAVY, True)]])
        text(s, M+Inches(0.75), y+Inches(0.42), w-Inches(0.9), Inches(0.32), [[(d, 10, GRAY, False)]], line_spacing=1.0)
        y += rh+Inches(0.1)
    return y

def checklist(s, y, title, accent, items):
    if title:
        text(s, M, y, PW-2*M, Inches(0.4), [[(title, 14, NAVY, True)]]); y=y+Inches(0.45)
    for it in items:
        rect(s, M, y, PW-2*M, Inches(0.5), LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.12)
        rect(s, M+Inches(0.16), y+Inches(0.11), Inches(0.28), Inches(0.28), WHITE, line=accent, line_w=1.5, shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.2)
        text(s, M+Inches(0.16), y+Inches(0.1), Inches(0.28), Inches(0.28), [[("✓", 11, accent, True)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(s, M+Inches(0.62), y, PW-2*M-Inches(0.8), Inches(0.5), [[(it, 11.5, INK, False)]], anchor=MSO_ANCHOR.MIDDLE)
        y += Inches(0.58)
    return y

def ws_box(s, x, y, w, h, label, accent, nlines=0, hint=None):
    rect(s, x, y, w, h, WHITE, line=LINE, line_w=1.2, shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.04)
    rect(s, x, y, Inches(0.1), h, accent, shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.5)
    text(s, x+Inches(0.25), y+Inches(0.12), w-Inches(0.5), Inches(0.3), [[(label, 11.5, NAVY, True)]])
    if hint:
        text(s, x+Inches(0.25), y+Inches(0.4), w-Inches(0.5), Inches(0.3), [[(hint, 9.5, GRAY, False)]])
    if nlines:
        gap=(h-Inches(0.55))/nlines
        for i in range(nlines):
            yy=y+Inches(0.6)+gap*i
            rect(s, x+Inches(0.25), yy, w-Inches(0.5), Pt(1), LINE)

def tip(s, y, txt, accent=TEAL, label="TIP", h=Inches(0.55)):
    pale = {TEAL:PALE_TEAL, BLUE:PALE_BLUE, ORANGE:PALE_ORG, RED:PALE_RED}.get(accent, PALE_TEAL)
    rect(s, M, y, PW-2*M, h, pale, shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.08)
    text(s, M+Inches(0.22), y, PW-2*M-Inches(0.44), h, [[(label+"   ", 11, accent, True),(txt, 11, INK, False)]], anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.08)

def footer(s, txt, accent):
    rect(s, 0, PH-Inches(0.58), PW, Inches(0.58), NAVY)
    rect(s, 0, PH-Inches(0.58), PW, Inches(0.045), accent)
    text(s, M, PH-Inches(0.56), PW-2*M, Inches(0.56), [[(txt, 12.5, WHITE, True)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

def divider(s, part, kr, en, desc, accent, items):
    bg(s, NAVY)
    rect(s, 0, 0, PW, Inches(0.22), accent)
    rect(s, 0, Inches(0.22), PW, Inches(0.06), GOLD)
    rect(s, Inches(4.4), Inches(1.1), Inches(3.7), Inches(3.7), RGBColor(0x16,0x3A,0x66), shape=MSO_SHAPE.OVAL)
    rect(s, Inches(4.9), Inches(0.65), Inches(2.6), Inches(2.6), None, line=accent, line_w=1.5, shape=MSO_SHAPE.OVAL)
    text(s, M, Inches(2.1), Inches(4), Inches(0.6), [[(part, 17, accent, True)]])
    text(s, M, Inches(2.62), Inches(6.4), Inches(1.2), [[(kr, 34, WHITE, True)]])
    text(s, M, Inches(3.6), Inches(6.4), Inches(0.5), [[(en, 13.5, SUBSKY, True)]])
    rect(s, M, Inches(4.18), Inches(1.1), Inches(0.06), accent)
    text(s, M, Inches(4.42), Inches(6.3), Inches(1.0), [[(desc, 14, RGBColor(0xCF,0xDE,0xF0), False)]], line_spacing=1.3)
    y = Inches(5.95)
    for i,(t1,t2) in enumerate(items):
        rect(s, M, y, PW-2*M, Inches(1.0), CARDNAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.07)
        rect(s, Inches(0.72), y+Inches(0.24), Inches(0.52), Inches(0.52), accent, shape=MSO_SHAPE.OVAL)
        text(s, Inches(0.72), y+Inches(0.24), Inches(0.52), Inches(0.52), [[(str(i+1), 15, WHITE, True)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(s, Inches(1.5), y+Inches(0.18), Inches(5.4), Inches(0.4), [[(t1, 15, WHITE, True)]])
        text(s, Inches(1.5), y+Inches(0.56), Inches(5.4), Inches(0.35), [[(t2, 10.5, RGBColor(0xB7,0xCB,0xE3), False)]])
        y += Inches(1.12)

# ====================================================================
# 1. COVER
# ====================================================================
s = slide(); bg(s, NAVY)
rect(s, 0, 0, PW, Inches(0.22), TEAL)
rect(s, 0, Inches(0.22), PW, Inches(0.07), ORANGE)
rect(s, Inches(4.2), Inches(1.0), Inches(4.2), Inches(4.2), RGBColor(0x16,0x3A,0x66), shape=MSO_SHAPE.OVAL)
rect(s, Inches(5.0), Inches(0.5), Inches(2.3), Inches(2.3), RGBColor(0x1E,0x4A,0x80), shape=MSO_SHAPE.OVAL)
rect(s, Inches(4.55), Inches(1.35), Inches(3.5), Inches(3.5), None, line=TEAL, line_w=1.5, shape=MSO_SHAPE.OVAL)
text(s, M, Inches(1.5), Inches(5.0), Inches(0.5), [[("한전원자력연료  ·  신입직원 입문교육", 13, SUBSKY, True)]])
text(s, M, Inches(2.0), Inches(6.6), Inches(2.0), [[("생성형 AI 스마트워크", 33, WHITE, True)],[("& AI 활용 보고서 작성", 33, WHITE, True)]], line_spacing=1.05)
rect(s, M, Inches(3.55), Inches(1.2), Inches(0.07), ORANGE)
text(s, M, Inches(3.8), Inches(6.4), Inches(1.2),
     [[("제미나이 · 노트북LM · AI 스튜디오를 ", 14, RGBColor(0xCF,0xDE,0xF0), False),("모바일", 14, GOLD, True),("로 익히고,", 14, RGBColor(0xCF,0xDE,0xF0), False)],
      [("AI로 기획부터 보고서까지 — 교재만 보고 따라 하는 실습 가이드", 14, RGBColor(0xCF,0xDE,0xF0), False)]], line_spacing=1.3)
rect(s, M, Inches(7.4), PW-2*M, Inches(2.25), CARDNAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.04)
text(s, Inches(0.85), Inches(7.65), Inches(6), Inches(0.4), [[("교 육 개 요", 12, TEAL, True)]])
def info_row(y, k, v):
    text(s, Inches(0.85), y, Inches(1.4), Inches(0.35), [[(k, 12.5, GOLD, True)]])
    text(s, Inches(2.2), y, Inches(4.9), Inches(0.5), [[(v, 12.5, WHITE, False)]])
info_row(Inches(8.1), "일   시", "2026. 7. 8.(수)  13:00 ~ 18:00  (총 5H)")
info_row(Inches(8.55), "대   상", "신입직원 107명")
info_row(Inches(9.0), "구   성", "① 생성형 AI 스마트워크 (13~16시)")
text(s, Inches(2.2), Inches(9.4), Inches(4.9), Inches(0.5), [[("② AI 활용 보고서 작성 (16~18시)", 12.5, WHITE, False)]])
text(s, M, Inches(10.35), PW-2*M, Inches(0.4), [[("KEPCO Nuclear Fuel  ·  2026 신입직원 입문교육", 10.5, RGBColor(0x86,0xA6,0xCC), False)]], align=PP_ALIGN.CENTER)

# ====================================================================
# 2. OVERVIEW
# ====================================================================
s = slide(); bg(s)
header(s, "OVERVIEW  ·  과정 안내", "5시간의 여정", BLUE, TEAL, 2)
text(s, M, Inches(1.7), PW-2*M, Inches(0.6),
     [[("오후 한나절, ", 13, INK, False),("‘AI로 일하는 법’", 13, TEAL, True),("과 ", 13, INK, False),
       ("‘제대로 보고하는 법’", 13, ORANGE, True),("을 손으로 익힙니다.", 13, INK, False)]], line_spacing=1.2)
def journey(y, no, time, title, desc, accent, pale):
    rect(s, M, y, PW-2*M, Inches(1.4), pale, shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.06)
    rect(s, M, y, Inches(0.14), Inches(1.4), accent, shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.5)
    rect(s, Inches(0.78), y+Inches(0.27), Inches(0.88), Inches(0.88), accent, shape=MSO_SHAPE.OVAL)
    text(s, Inches(0.78), y+Inches(0.27), Inches(0.88), Inches(0.88), [[(no, 22, WHITE, True)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, Inches(1.95), y+Inches(0.2), Inches(4.9), Inches(0.4), [[(time, 11, accent, True)]])
    text(s, Inches(1.95), y+Inches(0.45), Inches(4.9), Inches(0.5), [[(title, 17, NAVY, True)]])
    text(s, Inches(1.95), y+Inches(0.92), Inches(4.9), Inches(0.4), [[(desc, 11, GRAY, False)]], line_spacing=1.1)
journey(Inches(2.4), "1", "PART 1  ·  13:00~16:00 (3H)", "생성형 AI 스마트워크", "제미나이 · 노트북LM · AI 스튜디오를 모바일로 직접 실습", TEAL, PALE_TEAL)
journey(Inches(4.0), "2", "PART 2  ·  16:00~18:00 (2H)", "AI 활용 보고서 작성", "보고의 기본·기획력 + 전지·매직으로 조별 보고서 작성", ORANGE, PALE_ORG)
rect(s, M, Inches(5.85), PW-2*M, Inches(1.95), LIGHT, line=LINE, line_w=1, shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.04)
text(s, Inches(0.8), Inches(6.05), Inches(6), Inches(0.4), [[("📱  이 교재 사용법", 14, NAVY, True)]])
for i,t1 in enumerate(["강의 흐름에 맞춰 핵심을 ‘읽고’, 빈칸·워크시트에 ‘적으며’ 따라오세요.",
                       "모든 실습은 개인 스마트폰으로 진행합니다 — PC가 없어도 OK.",
                       "교재의 프롬프트 예시는 그대로 복사해 사용해도 좋습니다.",
                       "오늘 만든 결과물이 곧 내일 업무에서 쓰는 ‘나의 도구’가 됩니다."]):
    text(s, Inches(0.85), Inches(6.5)+Inches(0.32)*i, Inches(5.9), Inches(0.35), [[("· ", 12, TEAL, True),(t1, 11.5, INK, False)]])
rect(s, M, Inches(8.05), PW-2*M, Inches(2.2), NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.04)
text(s, Inches(0.8), Inches(8.25), Inches(6), Inches(0.4), [[("준비물 체크", 12.5, GOLD, True)]])
for i,(a,b) in enumerate([("구글 계정","Gmail 로그인\n(없으면 즉시 가입)"),("크롬 / 사파리","모바일 웹브라우저"),("필기구","워크시트 작성용")]):
    xx = Inches(0.78)+Inches(2.08)*i
    rect(s, xx, Inches(8.75), Inches(1.92), Inches(1.25), CARDNAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.08)
    text(s, xx, Inches(8.92), Inches(1.92), Inches(0.4), [[(a, 12.5, WHITE, True)]], align=PP_ALIGN.CENTER)
    text(s, xx+Inches(0.1), Inches(9.32), Inches(1.72), Inches(0.7), [[(line, 9.5, RGBColor(0xBF,0xD2,0xE8), False)] for line in b.split("\n")], align=PP_ALIGN.CENTER, line_spacing=1.05)

# ====================================================================
# 3. PART 1 DIVIDER
# ====================================================================
s = slide()
divider(s, "PART 1  ·  13:00 ~ 16:00", "생성형 AI 스마트워크", "SMART WORK WITH GENERATIVE AI",
        "스마트폰만 있으면 누구나, 지금 바로. 세 가지 AI 도구로 업무 속도를 바꿉니다.",
        TEAL,
        [("제미나이 (Gemini)", "이미지·영상·문서를 만드는 만능 AI 어시스턴트"),
         ("노트북LM (NotebookLM)", "내 문서만 읽고 답하는 출처 기반 AI 분석가"),
         ("AI 스튜디오 (AI Studio)", "코딩 없이 챗봇·앱을 만드는 AI 개발 도구")])

# ====================================================================
# 4. 제미나이란 + 모바일 시작
# ====================================================================
s = slide(); bg(s)
header(s, "PART 1  제미나이", "제미나이란? & 모바일 시작하기", TEAL, TEAL, 4)
y = def_box(s, Inches(1.75), "한 줄 정의",
            "구글이 만든 AI 어시스턴트. 말이나 글로 요청하면 이미지·문서·번역·요약까지 만들어 줘요. 구글 계정만 있으면 무료!", TEAL, PALE_TEAL)
text(s, M, y+Inches(0.18), PW-2*M, Inches(0.4), [[("4가지 핵심 기능", 13.5, NAVY, True)]])
feats = [("🖼️","이미지","아이디어 스케치·시각자료"),("🎬","영상·음악","홍보 스크립트·사내방송 기획"),
         ("📝","캔버스","기획서·보고서·번역 전용 공간"),("💬","대화·요약","자료 요약·아이디어 정리")]
fy = y+Inches(0.6)
cw = (PW-2*M-Inches(0.36))/4
for i,(ic,t1,t2) in enumerate(feats):
    xx = M + (cw+Inches(0.12))*i
    rect(s, xx, fy, cw, Inches(1.5), WHITE, line=LINE, line_w=1, shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.07, shadow=True)
    text(s, xx, fy+Inches(0.18), cw, Inches(0.5), [[(ic, 22, INK, False)]], align=PP_ALIGN.CENTER)
    text(s, xx, fy+Inches(0.72), cw, Inches(0.3), [[(t1, 12, TEAL, True)]], align=PP_ALIGN.CENTER)
    text(s, xx+Inches(0.05), fy+Inches(1.02), cw-Inches(0.1), Inches(0.45), [[(t2, 8.5, GRAY, False)]], align=PP_ALIGN.CENTER, line_spacing=1.0)
my = fy+Inches(1.75)
steps(s, my, "모바일로 시작하기 — 3단계", TEAL,
      [("크롬·사파리에서 gemini.google.com 접속", "주소창에 직접 입력하거나 ‘제미나이’ 검색 → 첫 번째 링크"),
       ("구글 계정으로 로그인", "회사 계정 말고 개인 구글 계정 권장 · 별도 앱 설치 불필요"),
       ("대화창에 한국어로 요청 입력", "“○○ 보고서 초안 써줘”처럼 말하듯 입력하면 끝")])
tip(s, Inches(8.55), "안드로이드는 ‘Gemini 앱’, 아이폰은 ‘구글 앱’ 안에서도 바로 쓸 수 있어요.", TEAL, h=Inches(0.65))
footer(s, "말로 요청하면, AI가 일을 합니다", TEAL)

# ====================================================================
# 5. 제미나이 - 이미지/시각화
# ====================================================================
s = slide(); bg(s)
header(s, "PART 1  제미나이  |  기능 ①  이미지", "이미지로 아이디어를 빠르게", TEAL, TEAL, 5)
text(s, M, Inches(1.7), PW-2*M, Inches(0.5),
     [[("사용법  ", 11, TEAL, True),("대화창에 내용을 설명한 뒤 ‘이미지 만들기’ — 또는 “위 내용을 포스터로 그려줘”라고 직접 입력", 11, INK, False)]], line_spacing=1.1)
cw=(PW-2*M-Inches(0.2))/2
usage_card(s, M, Inches(2.35), cw, Inches(3.0), "1", "교육·홍보자료 시각화",
           "신입 안전교육 카드뉴스, 사내 캠페인 포스터 등 발표·게시용 이미지를 몇 초 만에 시안으로 받아볼 수 있어요.",
           "“원자력연료 품질의 중요성을 알리는 사내 안전 캠페인 포스터를 깔끔하고 신뢰감 있는 분위기로 그려줘.”", TEAL, PALE_TEAL)
usage_card(s, M+cw+Inches(0.2), Inches(2.35), cw, Inches(3.0), "2", "개념·공정 시각화",
           "글로만 된 설명을 보완하는 다이어그램·삽화 아이디어를 얻어 발표자료와 교육자료의 이해도를 높일 수 있어요.",
           "“핵연료 집합체의 기본 구조를 신입사원도 이해하기 쉽게 단순한 다이어그램 스타일로 그려줘. 레이블 포함.”", TEAL, PALE_TEAL)
steps(s, Inches(5.65), "이미지 만들기 — 3단계", TEAL,
      [("주제·분위기·용도를 구체적으로 입력", "“교육용 삽화”, “홍보 포스터”처럼 용도까지 적으면 결과가 좋아요"),
       ("‘이미지 만들기’ 버튼 또는 직접 요청", "“위 내용을 이미지로 그려줘” / “포스터 스타일로”"),
       ("결과 확인 후 수정 요청", "“더 밝게”, “가로 배치로” 처럼 추가로 다듬기")])
tip(s, Inches(8.95), "사진 속 부품·도표를 올리고 “이게 뭔지 설명해줘”처럼 이미지로 질문도 할 수 있어요.", TEAL)
footer(s, "찾지 말고, 만들어서 쓰세요", TEAL)

# ====================================================================
# 6. 제미나이 - 영상/음악 스크립트
# ====================================================================
s = slide(); bg(s)
header(s, "PART 1  제미나이  |  기능 ②  영상·음악", "영상 스크립트 & 사내방송 기획", TEAL, TEAL, 6)
cw=(PW-2*M-Inches(0.2))/2
usage_card(s, M, Inches(1.75), cw, Inches(3.0), "1", "홍보·교육 영상 스크립트",
           "사내 SNS·유튜브용 짧은 영상의 기획안과 대본을 순식간에 작성해 줘요. 촬영 전 구성을 빠르게 잡을 수 있어요.",
           "“신입사원이 소개하는 ‘우리 회사 1분 영상’ 쇼츠 스크립트를 써줘. 밝고 친근한 톤, 자막 문구 포함.”", TEAL, PALE_TEAL)
usage_card(s, M+cw+Inches(0.2), Inches(1.75), cw, Inches(3.0), "2", "행사·캠페인 음악 기획",
           "체육대회·안전주간 등 사내 행사 BGM 분위기와 응원가 가사 아이디어를 제안받아 기획 시간을 줄일 수 있어요.",
           "“전사 안전주간 캠페인에 어울리는 경쾌한 응원가 가사를 써줘. 후렴 반복, 안전 메시지 포함.”", TEAL, PALE_TEAL)
# 3rd usage full width
y=Inches(4.95)
rect(s, M, y, PW-2*M, Inches(1.7), WHITE, line=LINE, line_w=1, shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.03, shadow=True)
chip(s, M+Inches(0.2), y+Inches(0.18), Inches(0.95), Inches(0.3), "활용 3", TEAL, size=10)
text(s, M+Inches(1.28), y+Inches(0.16), Inches(5), Inches(0.34), [[("긴 영상·회의 자료 빠른 요약", 12.5, NAVY, True)]], anchor=MSO_ANCHOR.MIDDLE)
text(s, M+Inches(0.22), y+Inches(0.58), PW-2*M-Inches(0.44), Inches(0.5),
     [[("길이가 긴 교육 영상이나 세미나 자막의 핵심을 텍스트로 요약받아, 보고·전파 자료로 바로 활용할 수 있어요.", 11, INK, False)]], line_spacing=1.15)
rect(s, M+Inches(0.2), y+Inches(1.1), PW-2*M-Inches(0.4), Inches(0.45), LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.1)
text(s, M+Inches(0.35), y+Inches(1.1), PW-2*M-Inches(0.6), Inches(0.45),
     [[("💬  ", 10, TEAL, True),("“이 영상(자막)의 핵심을 5가지 요점과 키워드로 정리해줘.”", 10, GRAY, False)]], anchor=MSO_ANCHOR.MIDDLE)
tip(s, Inches(6.95), "가사·대본은 먼저 ‘초안’을 만들고 “더 짧게/더 신나게”처럼 다듬으면 품질이 올라가요.", TEAL)
# caution
y2=Inches(7.7)
rect(s, M, y2, PW-2*M, Inches(2.35), PALE_RED, shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.04)
rect(s, M, y2, Inches(0.13), Inches(2.35), RED, shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.5)
text(s, Inches(0.8), y2+Inches(0.2), Inches(6), Inches(0.4), [[("🔒  보안 — 꼭 지켜주세요 (원자력 업종 특성)", 13, RED, True)]])
for i,t1 in enumerate(["대외비·보안자료, 미공개 기술자료, 개인정보는 공개형 AI에 입력하지 않습니다.",
                       "민감 정보는 ‘일반화된 예시’로 바꿔 질문하세요. (실제 수치·도면 제외)",
                       "AI 결과물은 ‘초안’입니다 — 사실관계·기술 내용은 반드시 담당자·전문가가 검증.",
                       "회사 보안지침과 정보보호 규정을 항상 우선 적용합니다."]):
    text(s, Inches(0.82), y2+Inches(0.62)+Inches(0.4)*i, Inches(6.0), Inches(0.4),
         [[("· ", 12, RED, True),(t1, 11, INK, False)]], line_spacing=1.05)
footer(s, "민감정보는 빼고, 결과는 검증하고", TEAL)

# ====================================================================
# 7. 제미나이 - 캔버스
# ====================================================================
s = slide(); bg(s)
header(s, "PART 1  제미나이  |  기능 ③  캔버스", "캔버스 — 기획서 · 보고서 작성", TEAL, TEAL, 7)
y=def_box(s, Inches(1.75), "캔버스란?",
          "대화창 옆에 열리는 ‘문서 편집 공간’. 긴 글을 쓰고 단락별로 다듬는 문서 전용 작업실로, 기획·보고 업무에 특히 유용해요.", TEAL, PALE_TEAL)
y=steps(s, y+Inches(0.15), "캔버스 시작하기 — 3단계", TEAL,
        [("대화창에 문서 작성 요청", "“○○ 교육 운영 기획서 초안 써줘”처럼 종류·목적을 구체적으로"),
         ("캔버스 창에서 초안 확인", "제미나이가 문서를 옆 창에 작성 — 단락별로 검토"),
         ("단락 선택 후 수정 지시", "고칠 부분을 드래그 → “더 구체적으로” / “3줄로 줄여줘”")])
text(s, M, y+Inches(0.1), PW-2*M, Inches(0.4), [[("이렇게 활용하세요", 13.5, NAVY, True)]])
yy=y+Inches(0.52)
for ic,t1,t2 in [("📄","기획서·보고서 작성","교육·행사 기획서, 업무 보고서의 뼈대를 잡고 살을 붙여 완성"),
                 ("🎚️","글의 톤 바꾸기","딱딱한 보고서 → 쉬운 안내문 / 일반 글 → 격식 있는 공문체"),
                 ("🌐","다국어 번역·감수","해외 협력사 메일·문서를 영문으로 번역하고 자연스럽게 감수")]:
    rect(s, M, yy, PW-2*M, Inches(0.78), LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.06)
    text(s, M+Inches(0.2), yy, Inches(0.6), Inches(0.78), [[(ic, 18, INK, False)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, M+Inches(0.85), yy+Inches(0.12), Inches(5.5), Inches(0.32), [[(t1, 12.5, TEAL, True)]])
    text(s, M+Inches(0.85), yy+Inches(0.43), Inches(5.5), Inches(0.3), [[(t2, 10.5, GRAY, False)]])
    yy+=Inches(0.88)
tip(s, yy+Inches(0.05), "캔버스로 만든 보고서는 PART 2(보고서 작성)에서 그대로 이어서 활용합니다.", ORANGE)
footer(s, "AI로 ‘빈 화면의 공포’를 없앤다", TEAL)

# ====================================================================
# 8. 프롬프트 6원칙
# ====================================================================
s = slide(); bg(s)
header(s, "PART 1  제미나이", "프롬프트 6원칙  R·O·C·F·E·C", TEAL, TEAL, 8)
text(s, M, Inches(1.7), PW-2*M, Inches(0.55),
     [[("좋은 질문이 좋은 결과를 만듭니다. ", 12.5, INK, False),("이 6가지만 챙기면 결과 품질이 확 올라가요.", 12.5, TEAL, True)]], line_spacing=1.15)
rows=[("R","역할","“너는 원자력연료 회사의 교육 담당자야”",TEAL),
      ("O","목표","“신입 안전교육 기획안을 작성하고 싶어”",BLUE),
      ("C","맥락","“대상은 신입 107명, 장소는 사내 강당, 2시간 분량”",SKY),
      ("F","형식","“A4 1장, 목차·진행순서·준비물 항목 포함”",ORANGE),
      ("E","예시","“이런 형식(링크/사진)을 참고해서 만들어줘”",GOLD),
      ("C","제약","“전문 용어는 쉽게 풀어서, 한국어로만”",GRAY)]
yy=Inches(2.45)
for ab,t1,ex,col in rows:
    rect(s, M, yy, PW-2*M, Inches(0.92), WHITE, line=LINE, line_w=1, shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.05, shadow=True)
    rect(s, M+Inches(0.16), yy+Inches(0.16), Inches(0.6), Inches(0.6), col, shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.22)
    text(s, M+Inches(0.16), yy+Inches(0.14), Inches(0.6), Inches(0.6), [[(ab, 22, WHITE, True)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, M+Inches(0.95), yy+Inches(0.14), Inches(1.4), Inches(0.6), [[(t1, 14, NAVY, True)]], anchor=MSO_ANCHOR.MIDDLE)
    text(s, M+Inches(2.2), yy+Inches(0.14), Inches(4.2), Inches(0.64), [[(ex, 11, INK, False)]], anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.05)
    yy+=Inches(1.0)
tip(s, yy+Inches(0.02), "여섯 개를 한 문장에 다 넣을 필요는 없어요. 결과가 아쉬울 때 빠진 항목을 추가해 보세요.", TEAL)
footer(s, "막연하게 묻지 말고, 구체적으로 시켜라", TEAL)

# ====================================================================
# 9. 모바일 실습 ①
# ====================================================================
s = slide(); bg(s)
header(s, "PART 1  ·  모바일 실습 ①", "제미나이로 직접 해보기", TEAL, TEAL, 9)
rect(s, M, Inches(1.7), PW-2*M, Inches(0.7), PALE_TEAL, shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.06)
text(s, M+Inches(0.22), Inches(1.7), PW-2*M-Inches(0.44), Inches(0.7),
     [[("📱  스마트폰으로 gemini.google.com 접속 → 아래 미션을 순서대로 수행하고 결과를 메모하세요.", 11.5, INK, False)]], anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.1)
missions=[("MISSION 1","나를 소개하는 이미지 만들기","“신입사원인 나를 표현하는 밝고 전문적인 캐릭터 이미지를 그려줘.”"),
          ("MISSION 2","업무 메일 톤 바꾸기","아무 메일 문장을 붙여넣고 “정중한 비즈니스 말투로 다시 써줘.”"),
          ("MISSION 3","6원칙으로 기획안 요청","R·O·C·F·E·C를 넣어 “신입 워크숍 기획안”을 요청해보기")]
yy=Inches(2.6)
for tag,t1,ex in missions:
    rect(s, M, yy, PW-2*M, Inches(1.32), WHITE, line=LINE, line_w=1, shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.04, shadow=True)
    chip(s, M+Inches(0.2), yy+Inches(0.18), Inches(1.15), Inches(0.32), tag, TEAL, size=10)
    text(s, M+Inches(1.5), yy+Inches(0.16), Inches(5), Inches(0.36), [[(t1, 13, NAVY, True)]], anchor=MSO_ANCHOR.MIDDLE)
    rect(s, M+Inches(0.2), yy+Inches(0.62), PW-2*M-Inches(0.4), Inches(0.55), LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.08)
    text(s, M+Inches(0.38), yy+Inches(0.62), PW-2*M-Inches(0.7), Inches(0.55),
         [[("💬  ", 10, TEAL, True),(ex, 10.5, GRAY, False)]], anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.05)
    yy+=Inches(1.42)
ws_box(s, M, yy+Inches(0.08), PW-2*M, Inches(2.55), "✍️  가장 쓸모 있었던 결과 & 내 업무에 적용할 아이디어", TEAL, nlines=5)
footer(s, "써본 사람만 안다 — 일단 해보기", TEAL)

# ====================================================================
# 10. 노트북LM 이란 + 시작
# ====================================================================
s = slide(); bg(s)
header(s, "PART 1  노트북LM", "노트북LM이란? & 모바일 시작", TEAL, TEAL, 10)
y=def_box(s, Inches(1.75), "한 줄 정의",
          "내가 올린 문서(PDF·웹·유튜브 등)만 참고해 답하는 AI. 지어내기(할루시네이션)가 거의 없고, 답변마다 출처 페이지를 알려줘요.", TEAL, PALE_TEAL, h=Inches(1.2))
# value chips
text(s, M, y+Inches(0.12), PW-2*M, Inches(0.4), [[("왜 업무에 좋은가", 13, NAVY, True)]])
vy=y+Inches(0.55)
for i,(ic,t1,t2) in enumerate([("🎯","출처 기반","답변마다 근거 페이지 표시 — 신뢰도 ↑"),
                               ("🔒","내 자료만","공개 웹이 아닌 ‘내가 올린 자료’ 안에서만 답"),
                               ("⚡","요약·검색","수백 페이지를 몇 초 만에 핵심 정리")]):
    cw=(PW-2*M-Inches(0.3))/3
    xx=M+(cw+Inches(0.15))*i
    rect(s, xx, vy, cw, Inches(1.45), WHITE, line=LINE, line_w=1, shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.07, shadow=True)
    text(s, xx, vy+Inches(0.16), cw, Inches(0.5), [[(ic, 20, INK, False)]], align=PP_ALIGN.CENTER)
    text(s, xx, vy+Inches(0.66), cw, Inches(0.3), [[(t1, 12, TEAL, True)]], align=PP_ALIGN.CENTER)
    text(s, xx+Inches(0.08), vy+Inches(0.96), cw-Inches(0.16), Inches(0.45), [[(t2, 8.8, GRAY, False)]], align=PP_ALIGN.CENTER, line_spacing=1.0)
sy=vy+Inches(1.7)
steps(s, sy, "모바일로 시작하기", TEAL,
      [("notebooklm.google.com 접속 · 로그인", "구글 계정으로 바로 시작 — 완전 무료"),
       ("‘+ 새 노트북’ 만들고 자료 추가(Sources)", "PDF·구글문서·웹 URL·텍스트 붙여넣기 모두 가능"),
       ("인덱싱(1~3분) 후 질문 시작", "자연어로 질문하면 ‘출처 번호’가 붙은 답변이 나와요")])
tip(s, sy+Inches(2.95), "주제별로 노트북을 나눠 만드세요. 예: ‘신입교육’, ‘사내규정’, ‘안전매뉴얼’.", TEAL)
footer(s, "내 문서가 곧 나만의 AI 비서가 된다", TEAL)

# ====================================================================
# 11. 노트북LM 6대 기능
# ====================================================================
s = slide(); bg(s)
header(s, "PART 1  노트북LM", "이렇게 활용해요 — 6가지", TEAL, TEAL, 11)
text(s, M, Inches(1.7), PW-2*M, Inches(0.4), [[("규정·매뉴얼·보고서를 올려두고, 묻고 만들어 쓰세요.", 12, INK, False)]])
cards=[("01","자료 요약","긴 규정집·기술보고서의 핵심을 즉시 요약"),
       ("02","문서 비교","구버전 vs 신버전, 차이점을 표로 정리"),
       ("03","출처 Q&A","“관련 조항·페이지 찾아줘” 정확히 응답"),
       ("04","슬라이드 생성","올린 자료로 발표용 슬라이드 초안 자동 작성"),
       ("05","오디오 요약","문서를 팟캐스트형 음성 대화로 변환"),
       ("06","퀴즈·학습","핵심용어·OX퀴즈 등 교육자료 자동 생성")]
gy=Inches(2.25); cw=(PW-2*M-Inches(0.2))/2
for i,(no,t1,t2) in enumerate(cards):
    r=i//2; c=i%2
    xx=M+(cw+Inches(0.2))*c; yy=gy+(Inches(1.35))*r
    rect(s, xx, yy, cw, Inches(1.2), WHITE, line=LINE, line_w=1, shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.05, shadow=True)
    text(s, xx+Inches(0.2), yy+Inches(0.16), Inches(0.9), Inches(0.4), [[(no, 18, TEAL, True)]])
    rect(s, xx+Inches(0.2), yy+Inches(0.58), Inches(0.5), Inches(0.05), TEAL)
    text(s, xx+Inches(0.9), yy+Inches(0.14), cw-Inches(1.0), Inches(0.4), [[(t1, 13, NAVY, True)]], anchor=MSO_ANCHOR.MIDDLE)
    text(s, xx+Inches(0.22), yy+Inches(0.66), cw-Inches(0.4), Inches(0.5), [[(t2, 10, GRAY, False)]], line_spacing=1.1)
ey=gy+Inches(1.35)*3+Inches(0.05)
rect(s, M, ey, PW-2*M, Inches(1.55), PALE_TEAL, shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.04)
text(s, Inches(0.8), ey+Inches(0.18), Inches(6), Inches(0.4), [[("💼  업무 적용 예시", 13, TEAL, True)]])
for i,t1 in enumerate(["사내 규정집을 올려두고 “출장비 지급 기준 조항을 원문으로 찾아줘”",
                       "여러 회의록을 올려 “이번 분기 주요 결정사항만 요약해줘”",
                       "교육자료를 올려 “신입용 OX 퀴즈 10문제와 정답·해설 만들어줘”"]):
    text(s, Inches(0.82), ey+Inches(0.58)+Inches(0.3)*i, Inches(6), Inches(0.3), [[("· ", 11, TEAL, True),(t1, 10.5, INK, False)]])
footer(s, "지어내지 않는다, 출처를 보여준다", TEAL)

# ====================================================================
# 12. 노트북LM 실전 프롬프트 모음
# ====================================================================
s = slide(); bg(s)
header(s, "PART 1  노트북LM", "바로 쓰는 실전 프롬프트", TEAL, TEAL, 12)
text(s, M, Inches(1.7), PW-2*M, Inches(0.4), [[("자료를 올린 뒤, 그대로 복사해서 질문해 보세요.", 12, INK, False)]])
items=[("요약","이 문서의 핵심 내용을 5가지로 정리하고, 각각의 근거 페이지 번호도 알려줘."),
       ("비교","올린 두 보고서에서 ‘예산’과 ‘일정’ 부분을 비교해 차이점을 표로 만들어줘."),
       ("검색","이 규정집에서 ‘초과근무 수당 지급 기준’ 관련 조항을 정확한 원문으로 발췌해줘."),
       ("퀴즈","이 교육자료로 신입사원용 OX 퀴즈 10문제와 정답·해설을 만들어줘."),
       ("발표","이 보고서를 임원 보고용 5분 발표자료로 요약하고, 예상 질문 5개도 준비해줘."),
       ("오디오","이 안내자료를 5분 분량의 음성 가이드 스크립트로 바꿔줘.")]
yy=Inches(2.25)
for tag,body in items:
    rect(s, M, yy, PW-2*M, Inches(1.05), WHITE, line=LINE, line_w=1, shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.05, shadow=True)
    rect(s, M, yy, Inches(1.15), Inches(1.05), PALE_TEAL, shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.05)
    text(s, M, yy, Inches(1.15), Inches(1.05), [[(tag, 13, TEAL, True)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, M+Inches(1.35), yy, PW-2*M-Inches(1.55), Inches(1.05), [[("“"+body+"”", 11.5, INK, False)]], anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.15)
    yy+=Inches(1.15)
footer(s, "좋은 질문 한 줄이 1시간을 아낀다", TEAL)

# ====================================================================
# 13. 모바일 실습 ②
# ====================================================================
s = slide(); bg(s)
header(s, "PART 1  ·  모바일 실습 ②", "노트북LM으로 직접 해보기", TEAL, TEAL, 13)
rect(s, M, Inches(1.7), PW-2*M, Inches(0.85), PALE_TEAL, shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.06)
text(s, M+Inches(0.22), Inches(1.7), PW-2*M-Inches(0.44), Inches(0.85),
     [[("📱  notebooklm.google.com 접속 → 새 노트북 생성 → 강사가 제공하는(또는 임의의) PDF·웹 자료 1개를 올리고 아래 미션 수행", 11.5, INK, False)]], anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.12)
missions=[("MISSION 1","핵심 5줄 요약 받기","“이 자료의 핵심을 5가지로, 출처 페이지와 함께 정리해줘.”"),
          ("MISSION 2","출처 기반 질문하기","자료 내용 중 궁금한 점을 묻고 ‘출처 번호’를 확인"),
          ("MISSION 3","퀴즈 or 슬라이드 생성","“OX 퀴즈 5문제” 또는 “발표 슬라이드 초안” 만들기")]
yy=Inches(2.75)
for tag,t1,ex in missions:
    rect(s, M, yy, PW-2*M, Inches(1.3), WHITE, line=LINE, line_w=1, shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.04, shadow=True)
    chip(s, M+Inches(0.2), yy+Inches(0.18), Inches(1.15), Inches(0.32), tag, TEAL, size=10)
    text(s, M+Inches(1.5), yy+Inches(0.16), Inches(5), Inches(0.36), [[(t1, 13, NAVY, True)]], anchor=MSO_ANCHOR.MIDDLE)
    rect(s, M+Inches(0.2), yy+Inches(0.6), PW-2*M-Inches(0.4), Inches(0.55), LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.08)
    text(s, M+Inches(0.38), yy+Inches(0.6), PW-2*M-Inches(0.7), Inches(0.55),
         [[("💬  ", 10, TEAL, True),(ex, 10.5, GRAY, False)]], anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.05)
    yy+=Inches(1.4)
ws_box(s, M, yy+Inches(0.05), PW-2*M, Inches(2.45), "✍️  노트북LM을 어떤 업무 자료에 써보고 싶나요?", TEAL, nlines=4)
footer(s, "검색하지 말고, 내 자료에게 물어보라", TEAL)

# ====================================================================
# 14. AI 스튜디오
# ====================================================================
s = slide(); bg(s)
header(s, "PART 1  AI 스튜디오", "AI 스튜디오란? & 활용", TEAL, TEAL, 14)
y=def_box(s, Inches(1.75), "한 줄 정의",
          "구글의 최신 AI로 맞춤형 챗봇·웹앱을 만드는 개발 플랫폼. 코딩 없이 ‘말로 설명’하면 완성돼요. (aistudio.google.com · 무료)", TEAL, PALE_TEAL, h=Inches(1.2))
text(s, M, y+Inches(0.12), PW-2*M, Inches(0.4), [[("핵심 기능", 13, NAVY, True)]])
fy=y+Inches(0.55)
cw=(PW-2*M-Inches(0.36))/4
for i,(t1,t2) in enumerate([("대용량 분석","수천 페이지 자료를 한 번에"),("멀티모달","이미지·영상·음성 동시 처리"),
                            ("챗봇 제작","우리 회사 전용 AI 챗봇"),("업무 연동","Drive·Sheets 자동화")]):
    xx=M+(cw+Inches(0.12))*i
    rect(s, xx, fy, cw, Inches(1.15), WHITE, line=LINE, line_w=1, shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.07, shadow=True)
    text(s, xx+Inches(0.05), fy+Inches(0.18), cw-Inches(0.1), Inches(0.4), [[(t1, 11.5, TEAL, True)]], align=PP_ALIGN.CENTER)
    text(s, xx+Inches(0.08), fy+Inches(0.58), cw-Inches(0.16), Inches(0.5), [[(t2, 8.8, GRAY, False)]], align=PP_ALIGN.CENTER, line_spacing=1.05)
uy=fy+Inches(1.4)
usage=[("🤖","맞춤형 챗봇 만들기","‘사내 규정 안내봇’, ‘신입 온보딩 도우미’처럼 성격을 정해 코딩 없이 챗봇 제작 — 사내 게시판·키오스크에 연결 가능."),
       ("📊","업무 데이터 자동 분석","구글 시트의 설문·현황 데이터를 AI가 읽고 분석·요약 — 매월 수작업하던 정리를 자동화."),
       ("🗂️","대규모 자료 교차 분석","수백 페이지 규정·매뉴얼을 한 번에 올려 필요한 조항을 즉시 검색·정리.")]
for ic,t1,t2 in usage:
    rect(s, M, uy, PW-2*M, Inches(1.05), LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.05)
    text(s, M+Inches(0.18), uy, Inches(0.7), Inches(1.05), [[(ic, 20, INK, False)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, M+Inches(0.95), uy+Inches(0.14), Inches(5.4), Inches(0.32), [[(t1, 12.5, TEAL, True)]])
    text(s, M+Inches(0.95), uy+Inches(0.46), Inches(5.4), Inches(0.5), [[(t2, 10, GRAY, False)]], line_spacing=1.12)
    uy+=Inches(1.15)
tip(s, uy+Inches(0.0), "심화 도구예요. 오늘은 ‘이런 게 가능하구나’ 감만 잡고, 관심 있으면 차차 익혀보세요.", TEAL)
footer(s, "코딩 없이, 말로 만드는 나만의 AI", TEAL)

# ====================================================================
# 15. PART 1 정리 & 체크리스트
# ====================================================================
s = slide(); bg(s)
header(s, "PART 1  마무리", "정리 & 실습 체크리스트", TEAL, TEAL, 15)
sumcards=[("G","제미나이","이미지·영상·캔버스로 콘텐츠와 문서를 빠르게 / 프롬프트 6원칙으로 품질 UP"),
          ("N","노트북LM","내 문서만 참고 — 지어내기 최소, 출처 자동 / 요약·비교·퀴즈·슬라이드까지"),
          ("A","AI 스튜디오","대용량 분석·맞춤 챗봇·업무 자동화 (코딩 불필요)")]
yy=Inches(1.75)
for ab,t1,t2 in sumcards:
    rect(s, M, yy, PW-2*M, Inches(1.15), WHITE, line=LINE, line_w=1, shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.05, shadow=True)
    rect(s, M+Inches(0.18), yy+Inches(0.22), Inches(0.7), Inches(0.7), TEAL, shape=MSO_SHAPE.OVAL)
    text(s, M+Inches(0.18), yy+Inches(0.2), Inches(0.7), Inches(0.7), [[(ab, 22, WHITE, True)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, M+Inches(1.05), yy+Inches(0.18), Inches(5.3), Inches(0.36), [[(t1, 14.5, NAVY, True)]])
    text(s, M+Inches(1.05), yy+Inches(0.56), Inches(5.3), Inches(0.5), [[(t2, 10.5, GRAY, False)]], line_spacing=1.12)
    yy+=Inches(1.25)
checklist(s, yy+Inches(0.05), "오늘의 실습 체크리스트", TEAL,
          ["제미나이에 로그인하고, 이미지 1장을 만들어 봤다",
           "프롬프트 6원칙(R·O·C·F·E·C)으로 요청해 봤다",
           "노트북LM에 자료를 올리고, 출처가 붙은 답을 받아 봤다",
           "노트북LM으로 요약 또는 퀴즈를 만들어 봤다",
           "보안 수칙(민감정보 미입력)을 이해했다"])
footer(s, "도구는 익혔다 — 이제 ‘보고’로", ORANGE)

# ====================================================================
# 16. PART 2 DIVIDER
# ====================================================================
s = slide()
divider(s, "PART 2  ·  16:00 ~ 18:00", "AI 활용 보고서 작성", "REPORT WRITING WITH AI",
        "기본기(기획력·보고 원칙)를 다지고, AI와 손(전지·매직)으로 ‘한 장 보고서’를 완성합니다.",
        ORANGE,
        [("보고의 기본 & 기획력", "왜·무엇을·어떻게 — 보고의 원칙과 두괄식"),
         ("공문서 · 원페이지 보고서", "행정업무편람 기반 작성법과 한 장 보고 구조"),
         ("AI 활용 + 전지·매직 실습", "AI로 초안 → 조별 전지에 매직으로 완성·발표")])

# ====================================================================
# 17. 왜 보고인가
# ====================================================================
s = slide(); bg(s)
header(s, "PART 2  보고의 기본", "왜, 보고인가?", ORANGE, ORANGE, 17)
y=def_box(s, Inches(1.78), "보고란 무엇인가",
          "보고는 ‘다 끝낸 뒤 완벽하게 하는 것’이 아니라, 내 일의 상황을 상사와 ‘동기화’해 제때 판단을 돕는 행위입니다.", ORANGE, PALE_ORG, h=Inches(1.3))
cw=(PW-2*M-Inches(0.2))/2; cy=y+Inches(0.28)
ch=Inches(2.15)
rect(s, M, cy, cw, ch, PALE_RED, shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.05)
text(s, M+Inches(0.22), cy+Inches(0.2), cw-Inches(0.44), Inches(0.4), [[("❌  흔한 오해", 13, RED, True)]])
text(s, M+Inches(0.22), cy+Inches(0.7), cw-Inches(0.44), Inches(1.3),
     [[("“다 끝내고 완벽하게 보고해야지.”", 11.5, INK, True)],[("→ 그 사이 상사는 깜깜이. 불안이 쌓이고, 문제는 손쓸 수 없을 때 드러난다.", 11, GRAY, False)]], line_spacing=1.2, space_after=6)
xx=M+cw+Inches(0.2)
rect(s, xx, cy, cw, ch, PALE_TEAL, shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.05)
text(s, xx+Inches(0.22), cy+Inches(0.2), cw-Inches(0.44), Inches(0.4), [[("⭕  진짜 보고", 13, TEAL, True)]])
text(s, xx+Inches(0.22), cy+Inches(0.7), cw-Inches(0.44), Inches(1.3),
     [[("상황을 상사와 ‘공유’하는 것.", 11.5, INK, True)],[("→ 불안을 덜어주고 의사결정을 돕는다. 신입이 가장 빨리 신뢰를 얻는 방법.", 11, GRAY, False)]], line_spacing=1.2, space_after=6)
# 기획력
py=cy+ch+Inches(0.35)
rect(s, M, py, PW-2*M, Inches(3.0), LIGHT, line=LINE, line_w=1, shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.04)
text(s, Inches(0.8), py+Inches(0.22), Inches(6), Inches(0.4), [[("💡  보고의 출발점은 ‘기획력’", 14, ORANGE, True)]])
text(s, Inches(0.8), py+Inches(0.66), Inches(6), Inches(0.4),
     [[("좋은 보고서는 ‘쓰기’ 전에 ‘생각’에서 결정됩니다. 쓰기 전 3가지를 먼저 정하세요.", 11.5, INK, False)]], line_spacing=1.1)
for i,(t1,t2) in enumerate([("Why  목적","이 보고로 무엇을 결정하게 할 것인가?"),
                            ("Who  대상","누가 읽는가? 그가 궁금한 것은 무엇인가?"),
                            ("What  핵심","한 문장으로 줄이면? = 가장 하고 싶은 말")]):
    ly=py+Inches(1.28)+Inches(0.55)*i
    rect(s, Inches(0.8), ly, PW-Inches(1.6), Inches(0.46), WHITE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.18)
    text(s, Inches(1.0), ly, Inches(1.7), Inches(0.46), [[(t1, 11.5, ORANGE, True)]], anchor=MSO_ANCHOR.MIDDLE)
    text(s, Inches(2.6), ly, Inches(4.1), Inches(0.46), [[(t2, 11, INK, False)]], anchor=MSO_ANCHOR.MIDDLE)
footer(s, "작은 보고가 쌓여 ‘신뢰’가 된다", ORANGE)

# ====================================================================
# 18. 3원칙 + PREP
# ====================================================================
s = slide(); bg(s)
header(s, "PART 2  보고의 기본", "좋은 보고의 3원칙 & 두괄식", ORANGE, ORANGE, 18)
text(s, M, Inches(1.7), PW-2*M, Inches(0.4), [[("잘된 보고는 언제나 이 세 가지를 만족합니다.", 12, INK, False)]])
prin=[("適","적시성","필요한 순간에 — 늦은 100점보다 빠른 70점", ORANGE),
      ("正","정확성","사실과 의견을 구분해서", BLUE),
      ("簡","간결성","결론부터, 짧게", TEAL)]
cw=(PW-2*M-Inches(0.3))/3
for i,(han,t1,t2,col) in enumerate(prin):
    xx=M+(cw+Inches(0.15))*i
    rect(s, xx, Inches(2.25), cw, Inches(1.7), WHITE, line=LINE, line_w=1, shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.06, shadow=True)
    rect(s, xx+(cw-Inches(0.7))/2, Inches(2.42), Inches(0.7), Inches(0.7), col, shape=MSO_SHAPE.OVAL)
    text(s, xx+(cw-Inches(0.7))/2, Inches(2.4), Inches(0.7), Inches(0.7), [[(han, 22, WHITE, True)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, xx, Inches(3.2), cw, Inches(0.35), [[(t1, 13.5, NAVY, True)]], align=PP_ALIGN.CENTER)
    text(s, xx+Inches(0.1), Inches(3.55), cw-Inches(0.2), Inches(0.4), [[(t2, 9.5, GRAY, False)]], align=PP_ALIGN.CENTER, line_spacing=1.05)
# PREP
py=Inches(4.2)
text(s, M, py, PW-2*M, Inches(0.4), [[("두괄식 = 결론부터, 그리고 ", 13.5, NAVY, True),("PREP", 13.5, ORANGE, True)]])
rect(s, M, py+Inches(0.45), PW-2*M, Inches(1.55), LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.04)
prep=[("P","Point","결론부터"),("R","Reason","근거·이유"),("E","Example","사례·수치"),("P","Point","결론 재확인")]
cw2=(PW-2*M-Inches(0.6))/4
for i,(ab,t1,t2) in enumerate(prep):
    xx=M+Inches(0.15)+(cw2+Inches(0.1))*i
    rect(s, xx, py+Inches(0.62), cw2, Inches(1.2), WHITE, line=ORANGE, line_w=1.2, shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.08)
    text(s, xx, py+Inches(0.74), cw2, Inches(0.45), [[(ab, 24, ORANGE, True)]], align=PP_ALIGN.CENTER)
    text(s, xx, py+Inches(1.2), cw2, Inches(0.3), [[(t1, 10.5, NAVY, True)]], align=PP_ALIGN.CENTER)
    text(s, xx, py+Inches(1.48), cw2, Inches(0.3), [[(t2, 9.5, GRAY, False)]], align=PP_ALIGN.CENTER)
# example
ey=py+Inches(2.2)
rect(s, M, ey, PW-2*M, Inches(2.0), PALE_ORG, shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.04)
text(s, Inches(0.8), ey+Inches(0.16), Inches(6), Inches(0.35), [[("🌀 미괄식 (신입 디폴트)", 11.5, GRAY, True)]])
text(s, Inches(0.8), ey+Inches(0.5), Inches(6), Inches(0.5),
     [[("“어제 거래처 연락이 왔는데, 담당자가 바뀌어서… 결국 일정이 2일 밀릴 것 같습니다.”", 10.5, INK, False)]], line_spacing=1.1)
text(s, Inches(0.8), ey+Inches(1.05), Inches(6), Inches(0.35), [[("⭐ 두괄식 (일잘러)", 11.5, ORANGE, True)]])
text(s, Inches(0.8), ey+Inches(1.4), Inches(6), Inches(0.5),
     [[("“일정이 2일 지연될 것 같습니다.(결론) 담당자 교체 때문이고, 대안은 △△입니다.”", 10.5, INK, True)]], line_spacing=1.1)
footer(s, "첫 문장이 보고의 승부처다", ORANGE)

# ====================================================================
# 19. 공문서 작성법
# ====================================================================
s = slide(); bg(s)
header(s, "PART 2  기본 이론", "공문서 작성법 (행정업무편람)", ORANGE, ORANGE, 19)
y=def_box(s, Inches(1.75), "공문서란?",
          "회사의 의사를 공식적으로 기록·전달하는 문서. ‘정해진 형식’과 ‘정확한 표현’이 생명입니다. 행정업무편람의 기본 원칙을 따릅니다.", ORANGE, PALE_ORG, h=Inches(1.2))
text(s, M, y+Inches(0.12), PW-2*M, Inches(0.4), [[("작성 5원칙", 13.5, NAVY, True)]])
rules=[("정확성","사실에 근거 · 오탈자·숫자 오류 없이"),
       ("간결성","한 문장 한 뜻 · 군더더기 없이 짧게"),
       ("명확성","애매한 표현 금지 · 누가 봐도 같은 의미"),
       ("성실·예의","정중한 표현 · 받는 사람 입장 고려"),
       ("형식 준수","두문-본문-결문 · 항목 기호(1. 가. 1) …)")]
ry=y+Inches(0.55)
for i,(t1,t2) in enumerate(rules):
    rect(s, M, ry, PW-2*M, Inches(0.6), LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.1)
    rect(s, M+Inches(0.18), ry+Inches(0.13), Inches(0.34), Inches(0.34), ORANGE, shape=MSO_SHAPE.OVAL)
    text(s, M+Inches(0.18), ry+Inches(0.11), Inches(0.34), Inches(0.34), [[(str(i+1), 12, WHITE, True)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, M+Inches(0.7), ry, Inches(1.6), Inches(0.6), [[(t1, 12.5, NAVY, True)]], anchor=MSO_ANCHOR.MIDDLE)
    text(s, M+Inches(2.3), ry, Inches(4.5), Inches(0.6), [[(t2, 11, GRAY, False)]], anchor=MSO_ANCHOR.MIDDLE)
    ry+=Inches(0.68)
# 항목기호
rect(s, M, ry+Inches(0.05), PW-2*M, Inches(1.35), NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.04)
text(s, Inches(0.8), ry+Inches(0.22), Inches(6), Inches(0.35), [[("항목 구분 순서 (행정업무편람)", 12, GOLD, True)]])
text(s, Inches(0.8), ry+Inches(0.58), Inches(6), Inches(0.7),
     [[("1.  →  가.  →  1)  →  가)  →  (1)  →  (가)  →  ①  →  ㉮", 14, WHITE, True)],
      [("둘째 항목부터는 상위 항목 위치에서 1자(2타)씩 들여쓰기", 10.5, SUBSKY, False)]], line_spacing=1.3)
footer(s, "형식이 신뢰를 만든다", ORANGE)

# ====================================================================
# 20. 원페이지 보고서
# ====================================================================
s = slide(); bg(s)
header(s, "PART 2  기본 이론", "원페이지 보고서 — 한 장으로", ORANGE, ORANGE, 20)
y=def_box(s, Inches(1.75), "왜 한 장인가",
          "바쁜 의사결정자는 ‘한 장’을 원합니다. 핵심만 구조화하면 한 장으로 충분 — 길이가 아니라 ‘구성’이 실력입니다.", ORANGE, PALE_ORG, h=Inches(1.15))
text(s, M, y+Inches(0.12), PW-2*M, Inches(0.4), [[("원페이지 기본 골격", 13.5, NAVY, True)]])
struct=[("제목","무엇에 대한 보고인가 — 한눈에 주제 파악"),
        ("개요 / 목적","왜 이 보고를 하는가 (배경·추진 근거)"),
        ("현황 / 문제","지금 상황은? 핵심 사실·데이터 (사실 중심)"),
        ("대안 / 방안","어떻게 할 것인가 — 1~3개 선택지와 비교"),
        ("결론 / 건의","제언과 의사결정 요청 (무엇을 승인받을지)"),
        ("향후 일정","언제·누가·무엇을 — 실행 계획")]
sy=y+Inches(0.55)
for i,(t1,t2) in enumerate(struct):
    rect(s, M, sy, PW-2*M, Inches(0.7), WHITE, line=LINE, line_w=1, shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.06)
    rect(s, M, sy, Inches(0.12), Inches(0.7), ORANGE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.5)
    text(s, M+Inches(0.3), sy, Inches(1.85), Inches(0.7), [[(t1, 12.5, ORANGE, True)]], anchor=MSO_ANCHOR.MIDDLE)
    text(s, M+Inches(2.2), sy, Inches(4.5), Inches(0.7), [[(t2, 10.8, INK, False)]], anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.05)
    sy+=Inches(0.78)
tip(s, sy+Inches(0.02), "‘사실’과 ‘의견(제언)’을 칸으로 분리하면 신뢰도가 올라갑니다.", ORANGE)
footer(s, "길게 쓰지 말고, 구조로 압축하라", ORANGE)

# ====================================================================
# 21. AI로 보고서 빠르게 - 워크플로우
# ====================================================================
s = slide(); bg(s)
header(s, "PART 2  AI 활용", "AI로 보고서 빠르게 — 4단계", ORANGE, ORANGE, 21)
text(s, M, Inches(1.7), PW-2*M, Inches(0.5),
     [[("AI는 ", 12, INK, False),("‘초안 작성·구조화·다듬기’", 12, ORANGE, True),
       ("를 돕습니다. ‘판단과 책임’은 나의 몫.", 12, INK, False)]], line_spacing=1.1)
flow=[("기획","목적·대상·핵심 메시지를 한 줄로 정리 (AI에게 줄 ‘재료’)",
       "“신입교육 만족도 개선 보고서를 쓸 거야. 대상은 팀장, 핵심은 ‘실습 비중 확대 건의’.”"),
      ("초안","AI에게 원페이지 구조로 초안 요청",
       "“위 내용을 원페이지 보고서로. 개요·현황·개선안·건의·일정 순서로 써줘.”"),
      ("검증","사실·수치·근거를 내가 직접 확인·수정",
       "AI가 지어낸 수치·내용이 없는지 점검 → 실제 데이터로 교체 (가장 중요!)"),
      ("다듬기","톤·길이·표현을 AI로 마무리",
       "“더 간결하게”, “두괄식으로”, “공문체로” 등으로 마지막 손질")]
fy=Inches(2.35)
for i,(t1,t2,ex) in enumerate(flow):
    h=Inches(1.55)
    rect(s, M, fy, PW-2*M, h, WHITE, line=LINE, line_w=1, shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.04, shadow=True)
    rect(s, M+Inches(0.16), fy+Inches(0.2), Inches(0.55), Inches(0.55), ORANGE, shape=MSO_SHAPE.OVAL)
    text(s, M+Inches(0.16), fy+Inches(0.2), Inches(0.55), Inches(0.55), [[(str(i+1), 16, WHITE, True)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, M+Inches(0.88), fy+Inches(0.16), Inches(1.5), Inches(0.4), [[(t1, 14, NAVY, True)]])
    text(s, M+Inches(0.88), fy+Inches(0.52), PW-2*M-Inches(1.1), Inches(0.4), [[(t2, 10.5, GRAY, False)]], line_spacing=1.0)
    rect(s, M+Inches(0.2), fy+Inches(0.92), PW-2*M-Inches(0.4), Inches(0.5), LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.08)
    text(s, M+Inches(0.36), fy+Inches(0.92), PW-2*M-Inches(0.7), Inches(0.5),
         [[("💬  ", 9.5, ORANGE, True),(ex, 9.5, INK, False)]], anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.05)
    fy+=Inches(1.65)
footer(s, "AI는 거들 뿐 — 판단은 사람이", ORANGE)

# ====================================================================
# 22. 전지+매직 실습 안내
# ====================================================================
s = slide(); bg(s)
header(s, "PART 2  ·  팀 실습", "전지 · 매직 보고서 작성", ORANGE, ORANGE, 22)
rect(s, M, Inches(1.7), PW-2*M, Inches(1.1), PALE_ORG, shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.05)
rect(s, M, Inches(1.7), Inches(0.13), Inches(1.1), ORANGE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.5)
text(s, Inches(0.8), Inches(1.84), Inches(6), Inches(0.4), [[("🎯  실습 목표", 13, ORANGE, True)]])
text(s, Inches(0.8), Inches(2.2), Inches(6.0), Inches(0.6),
     [[("스마트폰으로 AI 초안을 만들고, 조별로 ", 11.5, INK, False),("전지에 매직으로 ‘원페이지 보고서’", 11.5, ORANGE, True),
       ("를 완성해 발표합니다.", 11.5, INK, False)]], line_spacing=1.2)
steps(s, Inches(3.0), "진행 순서 (조별 · 약 40분)", ORANGE,
      [("주제 선정 & 기획", "미션카드에서 주제 선택 → 목적·대상·핵심 메시지 합의 (5분)"),
       ("AI로 초안 만들기 (모바일)", "제미나이·노트북LM으로 구조 초안 생성 → 사실 검증 (10분)"),
       ("전지에 옮겨 구성", "원페이지 골격에 맞춰 매직으로 작성 · 도식·강조 활용 (20분)"),
       ("조별 발표 & 피드백", "1분 두괄식 발표 → 상호 피드백 (5분)")])
# 평가기준
ey=Inches(7.05)
rect(s, M, ey, PW-2*M, Inches(2.2), NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.04)
text(s, Inches(0.8), ey+Inches(0.22), Inches(6), Inches(0.4), [[("🏆  평가 기준 (각 5점)", 13, GOLD, True)]])
crit=[("구성","원페이지 골격을 갖췄는가"),("두괄식","결론·핵심이 먼저 보이는가"),
      ("사실성","사실/의견 구분, 근거가 있는가"),("전달력","한눈에 이해되는 시각 구성인가")]
for i,(t1,t2) in enumerate(crit):
    r=i//2; c=i%2
    xx=Inches(0.8)+Inches(3.15)*c; yy=ey+Inches(0.78)+Inches(0.72)*r
    rect(s, xx, yy+Inches(0.06), Inches(0.1), Inches(0.5), ORANGE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.5)
    text(s, xx+Inches(0.22), yy, Inches(2.7), Inches(0.34), [[(t1, 12.5, WHITE, True)]])
    text(s, xx+Inches(0.22), yy+Inches(0.32), Inches(2.7), Inches(0.3), [[(t2, 9.8, SUBSKY, False)]], line_spacing=1.0)
footer(s, "손으로 쓰면, 머리에 남는다", ORANGE)

# ====================================================================
# 23. 실습 미션카드
# ====================================================================
s = slide(); bg(s)
header(s, "PART 2  ·  팀 실습", "보고서 미션카드 (택 1)", ORANGE, ORANGE, 23)
text(s, M, Inches(1.7), PW-2*M, Inches(0.4), [[("조별로 미션 1개를 골라 ‘원페이지 보고서’로 설계하세요. (실제 업무 상황으로 바꿔도 OK)", 11, INK, False)]], line_spacing=1.1)
cards=[("📋","MISSION A · 행사기획","신입사원 워크숍 개최 계획","목적·프로그램·예산·일정을 한 장으로 보고"),
       ("🛠️","MISSION B · 업무개선","반복 업무 자동화 제안","현황 문제 → AI 활용 개선안 → 기대효과 건의"),
       ("📊","MISSION C · 현황보고","교육 만족도 조사 결과","설문 결과 분석 → 시사점 → 개선 방안 제시"),
       ("⚠️","MISSION D · 상황보고","업무 지연 상황 중간보고","원인·영향·대응방안을 두괄식으로 신속 보고")]
gy=Inches(2.55); cw=(PW-2*M-Inches(0.2))/2; chh=Inches(2.75)
for i,(ic,tag,t1,t2) in enumerate(cards):
    r=i//2; c=i%2
    xx=M+(cw+Inches(0.2))*c; yy=gy+Inches(3.0)*r
    rect(s, xx, yy, cw, chh, WHITE, line=LINE, line_w=1, shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.05, shadow=True)
    rect(s, xx, yy, cw, Inches(0.56), PALE_ORG, shape=MSO_SHAPE.ROUND_2_SAME_RECTANGLE, round_=0.35)
    text(s, xx+Inches(0.22), yy+Inches(0.12), cw-Inches(0.44), Inches(0.34), [[(ic+"  "+tag, 11, ORANGE, True)]], anchor=MSO_ANCHOR.MIDDLE)
    text(s, xx+Inches(0.22), yy+Inches(0.82), cw-Inches(0.44), Inches(0.7), [[(t1, 14, NAVY, True)]], line_spacing=1.05)
    rect(s, xx+Inches(0.22), yy+Inches(1.62), cw-Inches(1.0), Inches(0.04), ORANGE)
    text(s, xx+Inches(0.22), yy+Inches(1.8), cw-Inches(0.44), Inches(0.8), [[(t2, 10.5, GRAY, False)]], line_spacing=1.2)
footer(s, "어떤 일이든, 보고는 한 장이면 된다", ORANGE)

# ====================================================================
# 24. 워크시트
# ====================================================================
s = slide(); bg(s)
header(s, "PART 2  ·  워크시트", "나의 원페이지 보고서 설계", ORANGE, ORANGE, 24)
text(s, M, Inches(1.68), PW-2*M, Inches(0.35), [[("선택한 미션을 아래 칸에 채우며 ‘전지 옮기기 전 설계도’를 완성하세요.", 10.5, GRAY, False)]])
ws_box(s, M, Inches(2.1), PW-2*M, Inches(0.85), "① 제목 / 한 줄 핵심 메시지 (가장 하고 싶은 말)", ORANGE, nlines=1)
cw=(PW-2*M-Inches(0.2))/2
ws_box(s, M, Inches(3.05), cw, Inches(1.3), "② 목적 (Why)", ORANGE, nlines=2, hint="이 보고로 결정할 것")
ws_box(s, M+cw+Inches(0.2), Inches(3.05), cw, Inches(1.3), "③ 대상 (Who)", ORANGE, nlines=2, hint="누가 읽는가")
ws_box(s, M, Inches(4.45), PW-2*M, Inches(1.5), "④ 현황 · 문제 (사실·데이터 중심)", ORANGE, nlines=3)
ws_box(s, M, Inches(6.05), PW-2*M, Inches(1.5), "⑤ 대안 · 방안 (어떻게 할 것인가)", ORANGE, nlines=3)
cw2=(PW-2*M-Inches(0.2))/2
ws_box(s, M, Inches(7.65), cw2, Inches(1.55), "⑥ 결론 · 건의", ORANGE, nlines=3, hint="무엇을 승인받을지")
ws_box(s, M+cw2+Inches(0.2), Inches(7.65), cw2, Inches(1.55), "⑦ 향후 일정", ORANGE, nlines=3, hint="언제·누가·무엇을")
tip(s, Inches(9.35), "막히면 스마트폰으로 “이 칸에 들어갈 내용을 제안해줘”라고 AI에게 물어보세요.", ORANGE)
footer(s, "설계가 끝나면, 전지로!", ORANGE)

# ====================================================================
# 25. WRAP-UP 핵심정리
# ====================================================================
s = slide(); bg(s)
header(s, "WRAP-UP  ·  핵심 정리", "오늘 챙겨갈 5가지", BLUE, GOLD, 25)
takeaways=[("AI는 ‘도구’ — 일을 대신하는 게 아니라 더 빠르게 돕는다", TEAL),
           ("좋은 질문(프롬프트)이 좋은 결과를 만든다 — R·O·C·F·E·C", TEAL),
           ("결과물은 ‘초안’ — 사실·수치는 반드시 내가 검증한다", ORANGE),
           ("보고는 두괄식 — 결론부터, 사실과 의견을 구분해서", ORANGE),
           ("보안 최우선 — 민감·대외비 정보는 공개형 AI에 넣지 않는다", RED)]
yy=Inches(1.85)
for i,(t1,col) in enumerate(takeaways):
    rect(s, M, yy, PW-2*M, Inches(1.05), WHITE, line=LINE, line_w=1, shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.06, shadow=True)
    rect(s, M, yy, Inches(0.14), Inches(1.05), col, shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.5)
    rect(s, M+Inches(0.32), yy+Inches(0.27), Inches(0.5), Inches(0.5), col, shape=MSO_SHAPE.OVAL)
    text(s, M+Inches(0.32), yy+Inches(0.25), Inches(0.5), Inches(0.5), [[(str(i+1), 16, WHITE, True)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, M+Inches(1.0), yy, PW-2*M-Inches(1.2), Inches(1.05), [[(t1, 12.5, INK, True)]], anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.15)
    yy+=Inches(1.15)
footer(s, "오늘 배운 한 가지를, 내일 업무에서", GOLD)

# ====================================================================
# 26. 액션플랜 / THANK YOU
# ====================================================================
s = slide(); bg(s, NAVY)
rect(s, 0, 0, PW, Inches(0.22), TEAL)
rect(s, 0, Inches(0.22), PW, Inches(0.07), ORANGE)
rect(s, Inches(4.6), Inches(0.9), Inches(3.6), Inches(3.6), RGBColor(0x16,0x3A,0x66), shape=MSO_SHAPE.OVAL)
text(s, M, Inches(1.4), Inches(6), Inches(0.5), [[("ACTION PLAN", 14, GOLD, True)]])
text(s, M, Inches(1.85), Inches(6.4), Inches(0.8), [[("내일부터 이렇게 써보겠습니다", 26, WHITE, True)]])
rect(s, M, Inches(2.7), Inches(1.1), Inches(0.06), ORANGE)
# 3 declarations
decl=[("AI 활용 1가지","예: 회의록 요약을 노트북LM으로 자동화한다"),
      ("보고 습관 1가지","예: 모든 보고는 결론부터 — 첫 문장에 답을 담는다"),
      ("이번 주 실천","예: 다음 주간보고를 AI 초안 → 검증으로 작성해 본다")]
yy=Inches(3.1)
for i,(t1,t2) in enumerate(decl):
    rect(s, M, yy, PW-2*M, Inches(1.25), CARDNAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.05)
    rect(s, M+Inches(0.2), yy+Inches(0.2), Inches(0.5), Inches(0.5), TEAL if i<1 else (ORANGE if i<2 else GOLD), shape=MSO_SHAPE.OVAL)
    text(s, M+Inches(0.2), yy+Inches(0.2), Inches(0.5), Inches(0.5), [[(str(i+1), 16, WHITE, True)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, M+Inches(0.9), yy+Inches(0.18), Inches(5.4), Inches(0.4), [[(t1, 14, WHITE, True)]])
    text(s, M+Inches(0.9), yy+Inches(0.58), Inches(5.4), Inches(0.5), [[(t2, 10.5, SUBSKY, False)]], line_spacing=1.05)
    rect(s, M+Inches(0.9), yy+Inches(1.02), Inches(5.3), Pt(1), RGBColor(0x3A,0x55,0x78))
    yy+=Inches(1.4)
text(s, M, Inches(8.5), PW-2*M, Inches(0.8), [[("Thank You", 30, WHITE, True)]], align=PP_ALIGN.CENTER)
text(s, M, Inches(9.25), PW-2*M, Inches(0.5), [[("오늘의 한 가지를, 내일의 첫 업무에서.", 13, SUBSKY, False)]], align=PP_ALIGN.CENTER)
text(s, M, Inches(10.2), PW-2*M, Inches(0.4), [[("한전원자력연료  ·  2026 신입직원 입문교육", 10.5, RGBColor(0x86,0xA6,0xCC), False)]], align=PP_ALIGN.CENTER)

OUT="한전원자력연료_신입직원_입문교육_생성형AI_보고서작성.pptx"
prs.save(OUT)
print("SAVED", OUT, "slides:", len(prs.slides._sldIdLst))
