# -*- coding: utf-8 -*-
"""한전원자력연료 신입직원 입문교육 — 강사용 교안 (16:9 가로형)"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

NAVY  = RGBColor(0x0F, 0x24, 0x44)
BLUE  = RGBColor(0x1B, 0x5E, 0xA8)
SKY   = RGBColor(0x3E, 0x8E, 0xDE)
TEAL  = RGBColor(0x12, 0xA1, 0x8E)
ORANGE= RGBColor(0xE8, 0x7A, 0x2B)
GOLD  = RGBColor(0xF2, 0xB1, 0x3C)
LIGHT = RGBColor(0xF3, 0xF6, 0xFB)
LINE  = RGBColor(0xD9, 0xE2, 0xEE)
INK   = RGBColor(0x1B, 0x27, 0x38)
GRAY  = RGBColor(0x5C, 0x6B, 0x80)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RED   = RGBColor(0xD2, 0x4B, 0x3E)
PALE_TEAL = RGBColor(0xE7, 0xF5, 0xF2)
PALE_ORG  = RGBColor(0xFC, 0xF0, 0xE4)
PALE_BLUE = RGBColor(0xE9, 0xF1, 0xFB)
PALE_RED  = RGBColor(0xFB, 0xEA, 0xE8)
CARDNAVY  = RGBColor(0x16, 0x33, 0x5C)
SUBSKY    = RGBColor(0x9F, 0xC7, 0xF0)
FONT = "맑은 고딕"

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
PW, PH = prs.slide_width, prs.slide_height
M = Inches(0.6)

def slide(): return prs.slides.add_slide(BLANK)

def _set_font(run, size, color, bold, font=FONT):
    run.font.size=Pt(size); run.font.color.rgb=color; run.font.bold=bold; run.font.name=font
    rPr=run._r.get_or_add_rPr()
    for tag in ('a:ea','a:cs'):
        e=rPr.find(qn(tag))
        if e is None: e=rPr.makeelement(qn(tag),{}); rPr.append(e)
        e.set('typeface',font)

def rect(s,x,y,w,h,fill,line=None,line_w=0,shape=MSO_SHAPE.RECTANGLE,shadow=False,round_=None):
    sp=s.shapes.add_shape(shape,x,y,w,h)
    if fill is None: sp.fill.background()
    else: sp.fill.solid(); sp.fill.fore_color.rgb=fill
    if line is None: sp.line.fill.background()
    else: sp.line.color.rgb=line; sp.line.width=Pt(line_w)
    sp.shadow.inherit=False
    if shadow:
        el=sp._element.spPr; ef=el.makeelement(qn('a:effectLst'),{})
        sh=el.makeelement(qn('a:outerShdw'),{'blurRad':'90000','dist':'40000','dir':'5400000','rotWithShape':'0'})
        clr=el.makeelement(qn('a:srgbClr'),{'val':'1B2738'}); al=el.makeelement(qn('a:alpha'),{'val':'20000'})
        clr.append(al); sh.append(clr); ef.append(sh); el.append(ef)
    if round_ is not None and shape in (MSO_SHAPE.ROUNDED_RECTANGLE,MSO_SHAPE.ROUND_2_SAME_RECTANGLE):
        try: sp.adjustments[0]=round_
        except: pass
    return sp

def text(s,x,y,w,h,paras,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP,space_after=3,line_spacing=1.0,wrap=True):
    tb=s.shapes.add_textbox(x,y,w,h); tf=tb.text_frame; tf.word_wrap=wrap; tf.vertical_anchor=anchor
    tf.margin_left=0; tf.margin_right=0; tf.margin_top=0; tf.margin_bottom=0
    first=True
    for para in paras:
        p=tf.paragraphs[0] if first else tf.add_paragraph(); first=False
        p.alignment=align; p.space_after=Pt(space_after); p.space_before=Pt(0); p.line_spacing=line_spacing
        for r in para:
            run=p.add_run(); run.text=r[0]; _set_font(run,r[1],r[2],r[3],r[4] if len(r)>4 else FONT)
    return tb

def bg(s,color=WHITE): rect(s,0,0,PW,PH,color)

def chip(s,x,y,w,h,label,color,txt=WHITE,size=12):
    rect(s,x,y,w,h,color,shape=MSO_SHAPE.ROUNDED_RECTANGLE,round_=0.5)
    text(s,x,y,w,h,[[(label,size,txt,True)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)

def header(s,tag,title,accent,color=None):
    page=len(prs.slides._sldIdLst)
    rect(s,0,0,PW,Inches(0.16),accent)
    text(s,M,Inches(0.46),Inches(10),Inches(0.32),[[(tag,13,color or accent,True)]])
    text(s,M,Inches(0.82),PW-2*M-Inches(0.6),Inches(0.7),[[(title,30,NAVY,True)]])
    rect(s,M,Inches(1.62),Inches(1.0),Inches(0.06),accent)
    text(s,PW-Inches(1.3),Inches(0.5),Inches(0.7),Inches(0.4),[[(str(page),15,GRAY,True)]],align=PP_ALIGN.RIGHT)

def footer(s,msg,accent):
    rect(s,0,PH-Inches(0.62),PW,Inches(0.62),NAVY)
    rect(s,0,PH-Inches(0.62),PW,Inches(0.05),accent)
    text(s,M,PH-Inches(0.6),PW-2*M,Inches(0.58),[[(msg,15,WHITE,True)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)

def divider(s,part,kr,en,desc,accent,items):
    bg(s,NAVY)
    rect(s,0,0,PW,Inches(0.22),accent); rect(s,0,Inches(0.22),PW,Inches(0.07),GOLD)
    rect(s,Inches(8.0),Inches(1.2),Inches(4.6),Inches(4.6),RGBColor(0x16,0x3A,0x66),shape=MSO_SHAPE.OVAL)
    rect(s,Inches(8.9),Inches(0.7),Inches(3.0),Inches(3.0),None,line=accent,line_w=1.5,shape=MSO_SHAPE.OVAL)
    text(s,M,Inches(1.9),Inches(7),Inches(0.5),[[(part,20,accent,True)]])
    text(s,M,Inches(2.5),Inches(7.5),Inches(1.1),[[(kr,46,WHITE,True)]])
    text(s,M,Inches(3.7),Inches(7),Inches(0.5),[[(en,16,SUBSKY,True)]])
    rect(s,M,Inches(4.35),Inches(1.2),Inches(0.07),accent)
    text(s,M,Inches(4.6),Inches(7),Inches(0.8),[[(desc,17,RGBColor(0xCF,0xDE,0xF0),False)]],line_spacing=1.3)
    # agenda chips right side bottom
    y=Inches(5.55);
    for i,(t1,t2) in enumerate(items):
        x=M+(Inches(4.0)+Inches(0.15))*i
        rect(s,x,y,Inches(4.0),Inches(1.25),CARDNAVY,shape=MSO_SHAPE.ROUNDED_RECTANGLE,round_=0.06)
        rect(s,x+Inches(0.22),y+Inches(0.28),Inches(0.55),Inches(0.55),accent,shape=MSO_SHAPE.OVAL)
        text(s,x+Inches(0.22),y+Inches(0.28),Inches(0.55),Inches(0.55),[[(str(i+1),18,WHITE,True)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
        text(s,x+Inches(0.95),y+Inches(0.24),Inches(2.9),Inches(0.4),[[(t1,15,WHITE,True)]])
        text(s,x+Inches(0.95),y+Inches(0.66),Inches(2.9),Inches(0.45),[[(t2,11,SUBSKY,False)]],line_spacing=1.0)

def def_box(s,x,y,w,h,label,body,accent,pale):
    rect(s,x,y,w,h,pale,shape=MSO_SHAPE.ROUNDED_RECTANGLE,round_=0.05)
    rect(s,x,y,Inches(0.16),h,accent,shape=MSO_SHAPE.ROUNDED_RECTANGLE,round_=0.5)
    text(s,x+Inches(0.4),y+Inches(0.22),w-Inches(0.7),Inches(0.4),[[(label,16,accent,True)]])
    text(s,x+Inches(0.4),y+Inches(0.68),w-Inches(0.7),h-Inches(0.8),[[(body,15.5,INK,False)]],line_spacing=1.25)

# ====================================================================
# 1. COVER
# ====================================================================
s=slide(); bg(s,NAVY)
rect(s,0,0,PW,Inches(0.24),TEAL); rect(s,0,Inches(0.24),PW,Inches(0.08),ORANGE)
rect(s,Inches(8.4),Inches(0.8),Inches(5.2),Inches(5.2),RGBColor(0x16,0x3A,0x66),shape=MSO_SHAPE.OVAL)
rect(s,Inches(9.6),Inches(0.3),Inches(2.8),Inches(2.8),RGBColor(0x1E,0x4A,0x80),shape=MSO_SHAPE.OVAL)
rect(s,Inches(8.9),Inches(1.3),Inches(4.2),Inches(4.2),None,line=TEAL,line_w=1.5,shape=MSO_SHAPE.OVAL)
text(s,M,Inches(1.5),Inches(8),Inches(0.5),[[("한전원자력연료  ·  신입직원 입문교육",17,SUBSKY,True)]])
text(s,M,Inches(2.15),Inches(8.2),Inches(2.0),[[("생성형 AI 스마트워크",44,WHITE,True)],[("& AI 활용 보고서 작성",44,WHITE,True)]],line_spacing=1.05)
rect(s,M,Inches(4.15),Inches(1.5),Inches(0.08),ORANGE)
text(s,M,Inches(4.45),Inches(8),Inches(0.6),
     [[("제미나이 · 노트북LM · AI 스튜디오를 ",18,RGBColor(0xCF,0xDE,0xF0),False),("모바일",18,GOLD,True),("로 익히고, AI로 보고서까지",18,RGBColor(0xCF,0xDE,0xF0),False)]])
# info row
iy=Inches(5.7)
for i,(k,v) in enumerate([("일시","2026. 7. 8.(수) 13:00~18:00 (5H)"),("대상","신입직원 107명"),("구성","① 생성형 AI  ② AI 보고서 작성")]):
    x=M+(Inches(4.0)+Inches(0.06))*i
    rect(s,x,iy,Inches(4.0),Inches(1.05),CARDNAVY,shape=MSO_SHAPE.ROUNDED_RECTANGLE,round_=0.06)
    text(s,x+Inches(0.3),iy+Inches(0.16),Inches(3.5),Inches(0.3),[[(k,13,GOLD,True)]])
    text(s,x+Inches(0.3),iy+Inches(0.5),Inches(3.5),Inches(0.5),[[(v,13,WHITE,False)]],line_spacing=1.0)

# ====================================================================
# 2. AGENDA
# ====================================================================
s=slide(); bg(s)
header(s,"OVERVIEW · 과정 안내","오늘의 5시간",BLUE,TEAL)
def part_card(x,no,time,title,desc,accent,pale):
    w=Inches(5.85)
    rect(s,x,Inches(2.0),w,Inches(2.5),pale,shape=MSO_SHAPE.ROUNDED_RECTANGLE,round_=0.04)
    rect(s,x,Inches(2.0),w,Inches(0.16),accent)
    rect(s,x+Inches(0.4),Inches(2.45),Inches(1.1),Inches(1.1),accent,shape=MSO_SHAPE.OVAL)
    text(s,x+Inches(0.4),Inches(2.45),Inches(1.1),Inches(1.1),[[(no,34,WHITE,True)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    text(s,x+Inches(1.8),Inches(2.5),w-Inches(2.0),Inches(0.35),[[(time,14,accent,True)]])
    text(s,x+Inches(1.8),Inches(2.85),w-Inches(2.0),Inches(0.5),[[(title,23,NAVY,True)]])
    text(s,x+Inches(0.45),Inches(3.75),w-Inches(0.8),Inches(0.6),[[(desc,14,GRAY,False)]],line_spacing=1.2)
part_card(M,"1","13:00~16:00 (3H)","생성형 AI 스마트워크","제미나이·노트북LM·AI 스튜디오를 모바일로 직접 실습",TEAL,PALE_TEAL)
part_card(Inches(6.9),"2","16:00~18:00 (2H)","AI 활용 보고서 작성","보고 기본·기획력 + 전지·매직으로 조별 보고서 작성·발표",ORANGE,PALE_ORG)
# bottom strip: 진행 방식
rect(s,M,Inches(4.9),PW-2*M,Inches(1.55),LIGHT,line=LINE,line_w=1,shape=MSO_SHAPE.ROUNDED_RECTANGLE,round_=0.03)
text(s,M+Inches(0.4),Inches(5.1),Inches(6),Inches(0.4),[[("📱  진행 방식",16,NAVY,True)]])
for i,(t1) in enumerate(["모든 실습은 스마트폰으로 — PC 없이 OK","교재의 프롬프트는 그대로 따라 해도 좋아요","오늘 만든 결과물 = 내일 쓰는 나의 도구"]):
    x=M+Inches(0.4)+Inches(3.95)*i
    text(s,x,Inches(5.6),Inches(3.8),Inches(0.7),[[("· ",14,TEAL,True),(t1,13.5,INK,False)]],line_spacing=1.1)
footer(s,"AI로 일하는 법 + 제대로 보고하는 법",TEAL)

# ====================================================================
# 3. PART 1 DIVIDER
# ====================================================================
s=slide()
divider(s,"PART 1 · 13:00~16:00","생성형 AI 스마트워크","SMART WORK WITH GENERATIVE AI",
        "스마트폰만 있으면 누구나, 지금 바로. 세 가지 AI 도구로 업무 속도를 바꿉니다.",TEAL,
        [("제미나이","만능 AI 어시스턴트"),("노트북LM","내 문서 전용 AI 분석가"),("AI 스튜디오","코딩 없는 AI 개발도구")])

# ====================================================================
# 4. 3 TOOLS OVERVIEW
# ====================================================================
s=slide(); bg(s)
header(s,"PART 1 · 도구 소개","오늘 배우는 세 가지 AI",TEAL)
tools=[("제미나이","Gemini","🤖","만능 AI 어시스턴트","이미지·영상·문서·웹앱을 말 한마디로",TEAL,PALE_TEAL),
       ("노트북LM","NotebookLM","📁","내 문서 전용 분석가","올린 자료만 보고 출처와 함께 답변",BLUE,PALE_BLUE),
       ("AI 스튜디오","AI Studio","🛠️","AI 개발 플랫폼","코딩 없이 챗봇·앱 제작·자동화",ORANGE,PALE_ORG)]
w=Inches(3.85)
for i,(t1,en,ic,sub,desc,accent,pale) in enumerate(tools):
    x=M+(w+Inches(0.3))*i
    rect(s,x,Inches(2.0),w,Inches(4.0),WHITE,line=LINE,line_w=1,shape=MSO_SHAPE.ROUNDED_RECTANGLE,round_=0.04,shadow=True)
    rect(s,x,Inches(2.0),w,Inches(1.2),accent,shape=MSO_SHAPE.ROUND_2_SAME_RECTANGLE,round_=0.08)
    text(s,x,Inches(2.18),w,Inches(0.6),[[(ic,34,WHITE,False)]],align=PP_ALIGN.CENTER)
    text(s,x,Inches(2.78),w,Inches(0.4),[[(en,13,RGBColor(0xE8,0xF3,0xF1),True)]],align=PP_ALIGN.CENTER)
    text(s,x,Inches(3.4),w,Inches(0.5),[[(t1,24,NAVY,True)]],align=PP_ALIGN.CENTER)
    text(s,x+Inches(0.3),Inches(4.05),w-Inches(0.6),Inches(0.4),[[(sub,15,accent,True)]],align=PP_ALIGN.CENTER)
    text(s,x+Inches(0.4),Inches(4.6),w-Inches(0.8),Inches(0.9),[[(desc,14,GRAY,False)]],align=PP_ALIGN.CENTER,line_spacing=1.25)
    text(s,x,Inches(5.55),w,Inches(0.35),[[("무료 · 구글 계정",12,GRAY,False)]],align=PP_ALIGN.CENTER)
footer(s,"세 도구를 손에 익히면, 일하는 속도가 달라진다",TEAL)

# ====================================================================
# 5. 제미나이 소개 & 시작
# ====================================================================
s=slide(); bg(s)
header(s,"PART 1 · 제미나이","제미나이 — 무엇이고, 어떻게 시작하나",TEAL)
def_box(s,M,Inches(1.95),Inches(6.0),Inches(1.7),"한 줄 정의",
        "구글이 만든 AI 어시스턴트. 말이나 글로 요청하면 이미지·문서·번역·요약까지 만들어 줘요. 구글 계정만 있으면 무료!",TEAL,PALE_TEAL)
# features mini
text(s,M,Inches(4.0),Inches(6),Inches(0.4),[[("핵심 기능",16,NAVY,True)]])
for i,(ic,t1) in enumerate([("🖼️","이미지"),("🎬","영상·음악"),("📝","캔버스"),("💬","요약·대화")]):
    x=M+Inches(1.45)*i
    rect(s,x,Inches(4.45),Inches(1.3),Inches(1.3),LIGHT,line=LINE,line_w=1,shape=MSO_SHAPE.ROUNDED_RECTANGLE,round_=0.08)
    text(s,x,Inches(4.6),Inches(1.3),Inches(0.5),[[(ic,24,INK,False)]],align=PP_ALIGN.CENTER)
    text(s,x,Inches(5.35),Inches(1.3),Inches(0.3),[[(t1,12.5,TEAL,True)]],align=PP_ALIGN.CENTER)
# right: 시작 3단계
rx=Inches(7.1)
rect(s,rx,Inches(1.95),Inches(5.6),Inches(4.05),NAVY,shape=MSO_SHAPE.ROUNDED_RECTANGLE,round_=0.04)
text(s,rx+Inches(0.45),Inches(2.2),Inches(5),Inches(0.4),[[("📱 모바일로 시작하기",17,GOLD,True)]])
for i,(t1,d) in enumerate([("gemini.google.com 접속","크롬·사파리 주소창에 입력 / 앱 설치 불필요"),
                           ("구글 계정으로 로그인","개인 구글 계정 권장"),
                           ("한국어로 요청 입력","“○○ 보고서 초안 써줘”처럼 말하듯이")]):
    y=Inches(2.75)+Inches(1.02)*i
    rect(s,rx+Inches(0.45),y,Inches(0.6),Inches(0.6),TEAL,shape=MSO_SHAPE.OVAL)
    text(s,rx+Inches(0.45),y,Inches(0.6),Inches(0.6),[[(str(i+1),18,WHITE,True)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    text(s,rx+Inches(1.25),y-Inches(0.02),Inches(4.2),Inches(0.4),[[(t1,16,WHITE,True)]])
    text(s,rx+Inches(1.25),y+Inches(0.38),Inches(4.2),Inches(0.45),[[(d,12,SUBSKY,False)]],line_spacing=1.0)
footer(s,"말로 요청하면, AI가 일을 한다",TEAL)

# ====================================================================
# 6. 제미나이 핵심 기능 + 활용
# ====================================================================
s=slide(); bg(s)
header(s,"PART 1 · 제미나이","제미나이로 이런 일을 합니다",TEAL)
cards=[("🖼️","이미지·시각화","교육·홍보 포스터, 개념 다이어그램을 몇 초 만에 시안으로","“품질 안전 캠페인 포스터를 신뢰감 있게 그려줘”"),
       ("🎬","영상·음악 기획","홍보·교육 영상 스크립트, 행사 응원가·BGM 아이디어","“회사 1분 소개 쇼츠 스크립트를 자막과 함께 써줘”"),
       ("📝","캔버스(문서)","기획서·보고서 초안 작성, 톤 변경, 다국어 번역","“○○ 교육 기획서 초안 써줘” → 단락별 수정")]
w=Inches(3.85)
for i,(ic,t1,desc,ex) in enumerate(cards):
    x=M+(w+Inches(0.3))*i
    rect(s,x,Inches(2.0),w,Inches(3.9),WHITE,line=LINE,line_w=1,shape=MSO_SHAPE.ROUNDED_RECTANGLE,round_=0.04,shadow=True)
    text(s,x+Inches(0.4),Inches(2.3),w-Inches(0.8),Inches(0.7),[[(ic,30,INK,False)]])
    text(s,x+Inches(0.4),Inches(3.1),w-Inches(0.8),Inches(0.5),[[(t1,19,TEAL,True)]])
    text(s,x+Inches(0.4),Inches(3.7),w-Inches(0.8),Inches(1.1),[[(desc,14.5,INK,False)]],line_spacing=1.25)
    rect(s,x+Inches(0.4),Inches(4.85),w-Inches(0.8),Inches(0.85),LIGHT,shape=MSO_SHAPE.ROUNDED_RECTANGLE,round_=0.08)
    text(s,x+Inches(0.6),Inches(4.85),w-Inches(1.2),Inches(0.85),[[("💬 ",11,TEAL,True),(ex,12,GRAY,False)]],anchor=MSO_ANCHOR.MIDDLE,line_spacing=1.1)
footer(s,"찾지 말고, 만들어서 쓰세요",TEAL)

# ====================================================================
# 7. 제미나이 웹앱 만들기
# ====================================================================
s=slide(); bg(s)
header(s,"PART 1 · 제미나이","코딩 없이 웹앱·게임 만들기",TEAL)
text(s,M,Inches(1.9),PW-2*M,Inches(0.5),[[("“○○ 앱 만들어줘” → 실제 작동하는 웹앱을 만들어 링크로 공유. ",16,INK,False),("간단한 것부터 업무용까지.",16,TEAL,True)]])
apps=[("🎮","점심 내기 게임","룰렛·사다리로 결정"),("☕","커피 취합 앱","주문 모아 자동 집계"),("🗳️","즉석 투표","회식 메뉴·날짜 결정"),
      ("✅","준비물 체크앱","출장·교육 체크리스트"),("🧮","업무 계산기","출장비·근무시간 환산"),("📊","미니 설문폼","만족도·의견 집계")]
w=Inches(3.85)
for i,(ic,t1,t2) in enumerate(apps):
    r=i//3; c=i%3
    x=M+(w+Inches(0.3))*c; y=Inches(2.55)+Inches(1.6)*r
    rect(s,x,y,w,Inches(1.4),WHITE,line=LINE,line_w=1,shape=MSO_SHAPE.ROUNDED_RECTANGLE,round_=0.06,shadow=True)
    rect(s,x+Inches(0.3),y+Inches(0.35),Inches(0.7),Inches(0.7),PALE_TEAL,shape=MSO_SHAPE.OVAL)
    text(s,x+Inches(0.3),y+Inches(0.35),Inches(0.7),Inches(0.7),[[(ic,22,INK,False)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    text(s,x+Inches(1.2),y+Inches(0.32),w-Inches(1.4),Inches(0.4),[[(t1,16,TEAL,True)]])
    text(s,x+Inches(1.2),y+Inches(0.74),w-Inches(1.4),Inches(0.45),[[(t2,12.5,GRAY,False)]])
footer(s,"아이디어만 있으면, 앱이 된다",TEAL)

# ====================================================================
# 8. 프롬프트 6원칙
# ====================================================================
s=slide(); bg(s)
header(s,"PART 1 · 제미나이","좋은 결과를 만드는 프롬프트 6원칙",TEAL)
text(s,M,Inches(1.95),PW-2*M,Inches(0.4),[[("R · O · C · F · E · C  — 막연히 묻지 말고, 구체적으로 시키세요.",16,INK,False)]])
rows=[("R","역할","“너는 회사 교육 담당자야”",TEAL),("O","목표","“신입 안전교육 기획안을 쓰고 싶어”",BLUE),
      ("C","맥락","“대상 신입 107명, 강당, 2시간”",SKY),("F","형식","“A4 1장, 순서·준비물 포함”",ORANGE),
      ("E","예시","“이런 형식을 참고해서”",GOLD),("C","제약","“전문용어는 쉽게, 한국어로만”",GRAY)]
w=Inches(5.95)
for i,(ab,t1,ex,col) in enumerate(rows):
    r=i//2; c=i%2
    x=M+(w+Inches(0.3))*c; y=Inches(2.5)+Inches(1.2)*r
    rect(s,x,y,w,Inches(1.0),WHITE,line=LINE,line_w=1,shape=MSO_SHAPE.ROUNDED_RECTANGLE,round_=0.05,shadow=True)
    rect(s,x+Inches(0.2),y+Inches(0.2),Inches(0.6),Inches(0.6),col,shape=MSO_SHAPE.ROUNDED_RECTANGLE,round_=0.22)
    text(s,x+Inches(0.2),y+Inches(0.2),Inches(0.6),Inches(0.6),[[(ab,22,WHITE,True)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    text(s,x+Inches(1.0),y+Inches(0.15),Inches(1.4),Inches(0.7),[[(t1,17,NAVY,True)]],anchor=MSO_ANCHOR.MIDDLE)
    text(s,x+Inches(2.2),y+Inches(0.15),w-Inches(2.4),Inches(0.7),[[(ex,13.5,INK,False)]],anchor=MSO_ANCHOR.MIDDLE,line_spacing=1.05)
footer(s,"질문이 좋아야, 답이 좋다",TEAL)

# ====================================================================
# 9. 실습 ①
# ====================================================================
s=slide(); bg(s)
header(s,"PART 1 · 모바일 실습 ①","제미나이로 직접 해보기",TEAL)
rect(s,M,Inches(1.95),PW-2*M,Inches(0.8),PALE_TEAL,shape=MSO_SHAPE.ROUNDED_RECTANGLE,round_=0.06)
text(s,M+Inches(0.4),Inches(1.95),PW-2*M-Inches(0.8),Inches(0.8),[[("📱  gemini.google.com 접속 → 아래 미션을 순서대로 해보고 결과를 옆 사람과 공유하세요.",16,INK,False)]],anchor=MSO_ANCHOR.MIDDLE)
ms=[("1","나를 소개하는 이미지 만들기","“신입사원인 나를 밝고 전문적인 캐릭터로 그려줘”"),
    ("2","업무 메일 톤 바꾸기","문장 붙여넣고 “정중한 비즈니스 말투로 다시 써줘”"),
    ("3","6원칙으로 기획안 요청","R·O·C·F·E·C를 넣어 “워크숍 기획안” 요청")]
w=Inches(3.85)
for i,(n,t1,ex) in enumerate(ms):
    x=M+(w+Inches(0.3))*i
    rect(s,x,Inches(3.05),w,Inches(2.8),WHITE,line=LINE,line_w=1,shape=MSO_SHAPE.ROUNDED_RECTANGLE,round_=0.04,shadow=True)
    rect(s,x+Inches(0.4),Inches(3.35),Inches(0.8),Inches(0.8),TEAL,shape=MSO_SHAPE.OVAL)
    text(s,x+Inches(0.4),Inches(3.35),Inches(0.8),Inches(0.8),[[(n,28,WHITE,True)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    text(s,x+Inches(0.4),Inches(4.35),w-Inches(0.8),Inches(0.8),[[(t1,16.5,NAVY,True)]],line_spacing=1.1)
    rect(s,x+Inches(0.4),Inches(5.05),w-Inches(0.8),Inches(0.65),LIGHT,shape=MSO_SHAPE.ROUNDED_RECTANGLE,round_=0.1)
    text(s,x+Inches(0.55),Inches(5.05),w-Inches(1.1),Inches(0.65),[[(ex,12,GRAY,False)]],anchor=MSO_ANCHOR.MIDDLE,line_spacing=1.05)
footer(s,"써본 사람만 안다 — 일단 해보기",TEAL)

# ====================================================================
# 10. 노트북LM 소개
# ====================================================================
s=slide(); bg(s)
header(s,"PART 1 · 노트북LM","노트북LM — 내 문서 전용 AI",BLUE)
def_box(s,M,Inches(1.95),Inches(6.0),Inches(2.0),"한 줄 정의",
        "내가 올린 문서(PDF·웹·유튜브)만 참고해 답하는 AI. 지어내기(할루시네이션)가 거의 없고, 답변마다 출처 페이지를 알려줘요.",BLUE,PALE_BLUE)
for i,(ic,t1,t2) in enumerate([("🎯","출처 기반","근거 페이지 표시"),("🔒","내 자료만","올린 자료 안에서만"),("⚡","요약·검색","수백 쪽을 몇 초로")]):
    x=M+Inches(2.0)*i
    rect(s,x,Inches(4.25),Inches(1.85),Inches(1.65),LIGHT,line=LINE,line_w=1,shape=MSO_SHAPE.ROUNDED_RECTANGLE,round_=0.07)
    text(s,x,Inches(4.45),Inches(1.85),Inches(0.5),[[(ic,24,INK,False)]],align=PP_ALIGN.CENTER)
    text(s,x,Inches(5.05),Inches(1.85),Inches(0.35),[[(t1,14,BLUE,True)]],align=PP_ALIGN.CENTER)
    text(s,x+Inches(0.1),Inches(5.45),Inches(1.65),Inches(0.4),[[(t2,11,GRAY,False)]],align=PP_ALIGN.CENTER)
rx=Inches(7.1)
rect(s,rx,Inches(1.95),Inches(5.6),Inches(3.95),NAVY,shape=MSO_SHAPE.ROUNDED_RECTANGLE,round_=0.04)
text(s,rx+Inches(0.45),Inches(2.2),Inches(5),Inches(0.4),[[("📱 시작하기",17,GOLD,True)]])
for i,(t1,d) in enumerate([("notebooklm.google.com 접속","구글 계정으로 로그인 · 무료"),
                           ("‘+ 새 노트북’ → 자료 추가","PDF·문서·웹 URL·텍스트"),
                           ("질문 시작","출처 번호가 붙은 답변 확인")]):
    y=Inches(2.7)+Inches(1.0)*i
    rect(s,rx+Inches(0.45),y,Inches(0.6),Inches(0.6),BLUE,shape=MSO_SHAPE.OVAL)
    text(s,rx+Inches(0.45),y,Inches(0.6),Inches(0.6),[[(str(i+1),18,WHITE,True)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    text(s,rx+Inches(1.25),y-Inches(0.02),Inches(4.2),Inches(0.4),[[(t1,15.5,WHITE,True)]])
    text(s,rx+Inches(1.25),y+Inches(0.36),Inches(4.2),Inches(0.4),[[(d,12,SUBSKY,False)]])
footer(s,"내 문서가 곧 나만의 AI 비서가 된다",BLUE)

# ====================================================================
# 11. 노트북LM 6기능
# ====================================================================
s=slide(); bg(s)
header(s,"PART 1 · 노트북LM","규정·매뉴얼·보고서를 올려두고, 묻고 만들기",BLUE)
cards=[("01","자료 요약","긴 규정·보고서 핵심 즉시 요약"),("02","문서 비교","구·신버전 차이를 표로 정리"),
       ("03","출처 Q&A","“관련 조항·페이지 찾아줘”"),("04","슬라이드 생성","발표용 슬라이드 초안 자동"),
       ("05","오디오 요약","팟캐스트형 음성 대화로 변환"),("06","퀴즈·학습","핵심용어·OX퀴즈 자동 생성")]
w=Inches(3.85)
for i,(no,t1,t2) in enumerate(cards):
    r=i//3; c=i%3
    x=M+(w+Inches(0.3))*c; y=Inches(2.05)+Inches(1.75)*r
    rect(s,x,y,w,Inches(1.55),WHITE,line=LINE,line_w=1,shape=MSO_SHAPE.ROUNDED_RECTANGLE,round_=0.05,shadow=True)
    text(s,x+Inches(0.35),y+Inches(0.25),Inches(1.0),Inches(0.5),[[(no,24,BLUE,True)]])
    text(s,x+Inches(1.25),y+Inches(0.22),w-Inches(1.5),Inches(0.5),[[(t1,17,NAVY,True)]],anchor=MSO_ANCHOR.MIDDLE)
    text(s,x+Inches(0.35),y+Inches(0.85),w-Inches(0.6),Inches(0.6),[[(t2,13,GRAY,False)]],line_spacing=1.15)
footer(s,"지어내지 않는다, 출처를 보여준다",BLUE)

# ====================================================================
# 12. 노트북LM 실전 프롬프트
# ====================================================================
s=slide(); bg(s)
header(s,"PART 1 · 노트북LM","바로 쓰는 실전 프롬프트",BLUE)
items=[("요약","핵심 내용을 5가지로 정리하고 근거 페이지도 알려줘"),
       ("비교","두 보고서의 ‘예산·일정’을 비교해 표로 만들어줘"),
       ("검색","이 규정집에서 ‘초과근무 수당 기준’을 원문으로 발췌해줘"),
       ("퀴즈","이 교육자료로 신입용 OX 퀴즈 10문제와 해설을 만들어줘"),
       ("발표","이 보고서를 5분 발표자료로 요약하고 예상 질문도 준비해줘"),
       ("오디오","이 안내자료를 5분 음성 가이드 스크립트로 바꿔줘")]
w=Inches(5.95)
for i,(tag,body) in enumerate(items):
    r=i//2; c=i%2
    x=M+(w+Inches(0.3))*c; y=Inches(2.0)+Inches(1.25)*r
    rect(s,x,y,w,Inches(1.05),WHITE,line=LINE,line_w=1,shape=MSO_SHAPE.ROUNDED_RECTANGLE,round_=0.05,shadow=True)
    rect(s,x,y,Inches(1.3),Inches(1.05),PALE_BLUE,shape=MSO_SHAPE.ROUNDED_RECTANGLE,round_=0.05)
    text(s,x,y,Inches(1.3),Inches(1.05),[[(tag,15,BLUE,True)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    text(s,x+Inches(1.55),y,w-Inches(1.75),Inches(1.05),[[("“"+body+"”",13.5,INK,False)]],anchor=MSO_ANCHOR.MIDDLE,line_spacing=1.15)
footer(s,"좋은 질문 한 줄이 1시간을 아낀다",BLUE)

# ====================================================================
# 13. 실습 ②
# ====================================================================
s=slide(); bg(s)
header(s,"PART 1 · 모바일 실습 ②","노트북LM으로 직접 해보기",BLUE)
rect(s,M,Inches(1.95),PW-2*M,Inches(0.8),PALE_BLUE,shape=MSO_SHAPE.ROUNDED_RECTANGLE,round_=0.06)
text(s,M+Inches(0.4),Inches(1.95),PW-2*M-Inches(0.8),Inches(0.8),[[("📱  새 노트북 생성 → 자료 1개 업로드 → 아래 미션 수행",16,INK,False)]],anchor=MSO_ANCHOR.MIDDLE)
ms=[("1","핵심 5줄 요약 받기","“핵심을 5가지로, 출처 페이지와 함께”"),
    ("2","출처 기반 질문하기","궁금한 점을 묻고 ‘출처 번호’ 확인"),
    ("3","퀴즈/슬라이드 생성","“OX 퀴즈 5문제” 또는 “슬라이드 초안”")]
w=Inches(3.85)
for i,(n,t1,ex) in enumerate(ms):
    x=M+(w+Inches(0.3))*i
    rect(s,x,Inches(3.05),w,Inches(2.8),WHITE,line=LINE,line_w=1,shape=MSO_SHAPE.ROUNDED_RECTANGLE,round_=0.04,shadow=True)
    rect(s,x+Inches(0.4),Inches(3.35),Inches(0.8),Inches(0.8),BLUE,shape=MSO_SHAPE.OVAL)
    text(s,x+Inches(0.4),Inches(3.35),Inches(0.8),Inches(0.8),[[(n,28,WHITE,True)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    text(s,x+Inches(0.4),Inches(4.35),w-Inches(0.8),Inches(0.7),[[(t1,16.5,NAVY,True)]],line_spacing=1.1)
    rect(s,x+Inches(0.4),Inches(5.05),w-Inches(0.8),Inches(0.65),LIGHT,shape=MSO_SHAPE.ROUNDED_RECTANGLE,round_=0.1)
    text(s,x+Inches(0.55),Inches(5.05),w-Inches(1.1),Inches(0.65),[[(ex,12,GRAY,False)]],anchor=MSO_ANCHOR.MIDDLE,line_spacing=1.05)
footer(s,"검색하지 말고, 내 자료에게 물어보라",BLUE)

# ====================================================================
# 14. 딥리서치
# ====================================================================
s=slide(); bg(s)
header(s,"PART 1 · 딥리서치","깊이 있는 자료조사 — 밖은 웹, 안은 내 자료",TEAL)
def drc(x,accent,pale,tag,what,hows,exs):
    w=Inches(5.95)
    rect(s,x,Inches(1.95),w,Inches(4.05),WHITE,line=LINE,line_w=1,shape=MSO_SHAPE.ROUNDED_RECTANGLE,round_=0.03,shadow=True)
    rect(s,x,Inches(1.95),w,Inches(0.62),pale,shape=MSO_SHAPE.ROUND_2_SAME_RECTANGLE,round_=0.2)
    text(s,x+Inches(0.35),Inches(1.95),w-Inches(0.7),Inches(0.62),[[(tag,16,accent,True)]],anchor=MSO_ANCHOR.MIDDLE)
    text(s,x+Inches(0.35),Inches(2.72),w-Inches(0.7),Inches(0.9),[[(what,13.5,INK,False)]],line_spacing=1.2)
    text(s,x+Inches(0.35),Inches(3.72),w-Inches(0.7),Inches(0.3),[[("이렇게 써요",13,accent,True)]])
    yy=Inches(4.05)
    for j,h in enumerate(hows):
        rect(s,x+Inches(0.35),yy+Inches(0.02),Inches(0.3),Inches(0.3),accent,shape=MSO_SHAPE.OVAL)
        text(s,x+Inches(0.35),yy+Inches(0.02),Inches(0.3),Inches(0.3),[[(str(j+1),10,WHITE,True)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
        text(s,x+Inches(0.78),yy,w-Inches(1.1),Inches(0.32),[[(h,12.5,INK,False)]],anchor=MSO_ANCHOR.MIDDLE)
        yy+=Inches(0.4)
    rect(s,x+Inches(0.35),Inches(5.62),w-Inches(0.7),Inches(0.32),pale,shape=MSO_SHAPE.ROUNDED_RECTANGLE,round_=0.2)
    text(s,x+Inches(0.5),Inches(5.6),w-Inches(1.0),Inches(0.36),[[("활용: ",11,accent,True),(exs,11,INK,False)]],anchor=MSO_ANCHOR.MIDDLE)
drc(M,TEAL,PALE_TEAL,"🌐  제미나이 딥리서치 (외부·웹)",
    "주제를 주면 AI가 수십~수백 개 웹사이트를 스스로 조사·교차검증해 출처가 달린 종합 보고서를 자동 작성.",
    ["‘Deep Research’ 선택 후 주제 입력","AI가 짠 ‘조사 계획’ 확인·수정","몇 분 뒤, 출처 포함 보고서 완성"],
    "원자력·에너지 기술/정책·해외 동향 조사")
drc(Inches(7.0),BLUE,PALE_BLUE,"📁  노트북LM (내부·내 자료)",
    "내가 올린 자료(규정·보고서·논문) 안에서만 교차분석·요약. 답변마다 출처 페이지가 표시돼 신뢰도가 높아요.",
    ["자료(소스) 업로드","자연어로 깊이 있게 질문","‘노트북 가이드’로 요약·마인드맵"],
    "사내 규정·매뉴얼·과거 보고서 정리")
footer(s,"넓게는 제미나이, 깊게는 노트북LM · 결과는 출처 확인 필수",TEAL)

# ====================================================================
# 15. AI 스튜디오
# ====================================================================
s=slide(); bg(s)
header(s,"PART 1 · AI 스튜디오","코딩 없이 챗봇·앱 만들기",ORANGE)
def_box(s,M,Inches(1.95),Inches(6.0),Inches(1.7),"한 줄 정의",
        "구글 최신 AI로 맞춤형 챗봇·웹앱을 만드는 개발 플랫폼. ‘말로 설명’하면 완성돼요. (aistudio.google.com · 무료)",ORANGE,PALE_ORG)
for i,(t1,t2) in enumerate([("대용량 분석","수천 쪽 한 번에"),("멀티모달","이미지·영상·음성"),("챗봇 제작","회사 전용 AI"),("업무 연동","Drive·Sheets 자동화")]):
    x=M+Inches(1.45)*i
    rect(s,x,Inches(4.0),Inches(1.3),Inches(1.9),LIGHT,line=LINE,line_w=1,shape=MSO_SHAPE.ROUNDED_RECTANGLE,round_=0.07)
    text(s,x+Inches(0.1),Inches(4.25),Inches(1.1),Inches(0.7),[[(t1,13,ORANGE,True)]],align=PP_ALIGN.CENTER,line_spacing=1.0)
    text(s,x+Inches(0.1),Inches(5.1),Inches(1.1),Inches(0.7),[[(t2,11,GRAY,False)]],align=PP_ALIGN.CENTER,line_spacing=1.05)
rx=Inches(7.1)
us=[("🤖","맞춤형 챗봇","‘규정 안내봇’·‘온보딩 도우미’를 코딩 없이 제작"),
    ("📊","데이터 자동 분석","시트의 설문·현황을 AI가 읽고 요약"),
    ("🗂️","대규모 자료 분석","수백 쪽 규정을 올려 즉시 검색·정리")]
for i,(ic,t1,t2) in enumerate(us):
    y=Inches(1.95)+Inches(1.35)*i
    rect(s,rx,y,Inches(5.6),Inches(1.2),WHITE,line=LINE,line_w=1,shape=MSO_SHAPE.ROUNDED_RECTANGLE,round_=0.05,shadow=True)
    text(s,rx+Inches(0.3),y,Inches(0.9),Inches(1.2),[[(ic,26,INK,False)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    text(s,rx+Inches(1.2),y+Inches(0.2),Inches(4.2),Inches(0.4),[[(t1,16,ORANGE,True)]])
    text(s,rx+Inches(1.2),y+Inches(0.62),Inches(4.2),Inches(0.5),[[(t2,12.5,GRAY,False)]],line_spacing=1.1)
footer(s,"심화 도구 — ‘이런 게 가능하다’ 감만 잡아도 충분",ORANGE)

# ====================================================================
# 16. PART 1 정리
# ====================================================================
s=slide(); bg(s)
header(s,"PART 1 · 마무리","Part 1 정리 & 체크리스트",TEAL)
sums=[("제미나이","이미지·영상·캔버스·웹앱 / 프롬프트 6원칙",TEAL),
      ("노트북LM","내 문서만 — 출처 자동 / 요약·비교·퀴즈·슬라이드",BLUE),
      ("AI 스튜디오","대용량 분석·챗봇·자동화 (코딩 불필요)",ORANGE)]
for i,(t1,t2,col) in enumerate(sums):
    y=Inches(1.95)+Inches(0.95)*i
    rect(s,M,y,Inches(6.0),Inches(0.82),WHITE,line=LINE,line_w=1,shape=MSO_SHAPE.ROUNDED_RECTANGLE,round_=0.06,shadow=True)
    rect(s,M,y,Inches(0.16),Inches(0.82),col,shape=MSO_SHAPE.ROUNDED_RECTANGLE,round_=0.5)
    text(s,M+Inches(0.4),y+Inches(0.13),Inches(2.0),Inches(0.55),[[(t1,17,NAVY,True)]],anchor=MSO_ANCHOR.MIDDLE)
    text(s,M+Inches(2.3),y,Inches(3.6),Inches(0.82),[[(t2,12.5,GRAY,False)]],anchor=MSO_ANCHOR.MIDDLE,line_spacing=1.05)
rx=Inches(7.1)
rect(s,rx,Inches(1.95),Inches(5.6),Inches(3.6),PALE_TEAL,shape=MSO_SHAPE.ROUNDED_RECTANGLE,round_=0.04)
text(s,rx+Inches(0.4),Inches(2.2),Inches(5),Inches(0.4),[[("✓ 오늘의 체크리스트",16,TEAL,True)]])
for i,t1 in enumerate(["제미나이 로그인 & 이미지 1장 만들기","프롬프트 6원칙으로 요청해 보기","노트북LM에 자료 올리고 출처 답변 받기","요약·퀴즈 만들어 보기","보안 수칙(민감정보 미입력) 이해"]):
    text(s,rx+Inches(0.45),Inches(2.7)+Inches(0.55)*i,Inches(4.8),Inches(0.5),[[("☐  ",14,TEAL,True),(t1,13.5,INK,False)]],anchor=MSO_ANCHOR.MIDDLE)
footer(s,"도구는 익혔다 — 이제 ‘보고’로",ORANGE)

# ====================================================================
# 17. PART 2 DIVIDER
# ====================================================================
s=slide()
divider(s,"PART 2 · 16:00~18:00","AI 활용 보고서 작성","REPORT WRITING WITH AI",
        "기본기(기획력·보고 원칙)를 다지고, AI와 손(전지·매직)으로 ‘한 장 보고서’를 완성합니다.",ORANGE,
        [("보고의 기본","왜·무엇을·어떻게"),("공문서·원페이지","양식과 구조"),("AI+전지 실습","조별 작성·발표")])

# ====================================================================
# 18. 왜 보고인가
# ====================================================================
s=slide(); bg(s)
header(s,"PART 2 · 보고의 기본","왜, 보고인가?",ORANGE)
def_box(s,M,Inches(1.95),PW-2*M,Inches(1.2),"보고란 무엇인가",
        "보고는 ‘다 끝낸 뒤 완벽하게 하는 것’이 아니라, 내 일의 상황을 상사와 ‘동기화’해 제때 판단을 돕는 행위입니다.",ORANGE,PALE_ORG)
w=Inches(5.95)
rect(s,M,Inches(3.4),w,Inches(1.4),PALE_RED,shape=MSO_SHAPE.ROUNDED_RECTANGLE,round_=0.05)
text(s,M+Inches(0.35),Inches(3.6),w-Inches(0.7),Inches(0.4),[[("❌  흔한 오해",15,RED,True)]])
text(s,M+Inches(0.35),Inches(4.05),w-Inches(0.7),Inches(0.7),[[("“다 끝내고 완벽하게 보고해야지” → 그 사이 상사는 깜깜이, 문제는 손쓸 수 없을 때 드러난다.",13,INK,False)]],line_spacing=1.2)
rect(s,Inches(7.0),Inches(3.4),w,Inches(1.4),PALE_TEAL,shape=MSO_SHAPE.ROUNDED_RECTANGLE,round_=0.05)
text(s,Inches(7.35),Inches(3.6),w-Inches(0.7),Inches(0.4),[[("⭕  진짜 보고",15,TEAL,True)]])
text(s,Inches(7.35),Inches(4.05),w-Inches(0.7),Inches(0.7),[[("상황을 ‘공유’하는 것 → 불안을 덜고 의사결정을 돕는다. 신입이 가장 빨리 신뢰를 얻는 법.",13,INK,False)]],line_spacing=1.2)
rect(s,M,Inches(5.0),PW-2*M,Inches(1.4),LIGHT,line=LINE,line_w=1,shape=MSO_SHAPE.ROUNDED_RECTANGLE,round_=0.04)
text(s,M+Inches(0.4),Inches(5.18),Inches(6),Inches(0.4),[[("💡  보고의 출발점은 ‘기획력’ — 쓰기 전에 3가지를 먼저 정하라",15,ORANGE,True)]])
for i,(t1,t2) in enumerate([("Why 목적","무엇을 결정하게?"),("Who 대상","누가 읽고, 뭘 궁금해하나?"),("What 핵심","한 문장으로 줄이면?")]):
    x=M+Inches(0.4)+Inches(4.0)*i
    text(s,x,Inches(5.68),Inches(3.8),Inches(0.6),[[(t1+"  ",13.5,ORANGE,True),(t2,13,INK,False)]],anchor=MSO_ANCHOR.MIDDLE)
footer(s,"작은 보고가 쌓여 ‘신뢰’가 된다",ORANGE)

# ====================================================================
# 19. 3원칙 + PREP
# ====================================================================
s=slide(); bg(s)
header(s,"PART 2 · 보고의 기본","좋은 보고의 3원칙 & 두괄식",ORANGE)
pr=[("適","적시성","늦은 100점보다 빠른 70점",ORANGE),("正","정확성","사실과 의견을 구분",BLUE),("簡","간결성","결론부터, 짧게",TEAL)]
w=Inches(3.85)
for i,(han,t1,t2,col) in enumerate(pr):
    x=M+(w+Inches(0.3))*i
    rect(s,x,Inches(1.95),w,Inches(1.7),WHITE,line=LINE,line_w=1,shape=MSO_SHAPE.ROUNDED_RECTANGLE,round_=0.06,shadow=True)
    rect(s,x+Inches(0.35),Inches(2.2),Inches(0.9),Inches(0.9),col,shape=MSO_SHAPE.OVAL)
    text(s,x+Inches(0.35),Inches(2.2),Inches(0.9),Inches(0.9),[[(han,26,WHITE,True)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    text(s,x+Inches(1.45),Inches(2.3),w-Inches(1.6),Inches(0.4),[[(t1,18,NAVY,True)]])
    text(s,x+Inches(1.45),Inches(2.78),w-Inches(1.6),Inches(0.6),[[(t2,12.5,GRAY,False)]],line_spacing=1.1)
text(s,M,Inches(3.9),Inches(6),Inches(0.4),[[("두괄식 = 결론부터, 그리고 ",16,NAVY,True),("PREP",16,ORANGE,True)]])
prep=[("P","결론부터"),("R","근거·이유"),("E","사례·수치"),("P","결론 재확인")]
for i,(ab,t1) in enumerate(prep):
    x=M+(Inches(2.9)+Inches(0.1))*i
    rect(s,x,Inches(4.4),Inches(2.9),Inches(1.0),PALE_ORG,shape=MSO_SHAPE.ROUNDED_RECTANGLE,round_=0.08)
    text(s,x+Inches(0.2),Inches(4.5),Inches(0.9),Inches(0.8),[[(ab,30,ORANGE,True)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    text(s,x+Inches(1.0),Inches(4.4),Inches(1.8),Inches(1.0),[[(t1,15,NAVY,True)]],anchor=MSO_ANCHOR.MIDDLE)
rect(s,M,Inches(5.6),PW-2*M,Inches(0.82),PALE_TEAL,shape=MSO_SHAPE.ROUNDED_RECTANGLE,round_=0.06)
text(s,M+Inches(0.4),Inches(5.6),PW-2*M-Inches(0.8),Inches(0.82),
     [[("⭐ 두괄식 예: ",13,TEAL,True),("“일정이 2일 지연될 것 같습니다(결론). 담당자 교체 때문이고, 대안은 △△입니다.”",13.5,INK,False)]],anchor=MSO_ANCHOR.MIDDLE)
footer(s,"첫 문장이 보고의 승부처다",ORANGE)

# ====================================================================
# 20. 공문서 작성법
# ====================================================================
s=slide(); bg(s)
header(s,"PART 2 · 기본 이론","공문서 작성법 (행정업무편람)",ORANGE)
def_box(s,M,Inches(1.95),Inches(6.0),Inches(1.4),"공문서란?",
        "회사의 의사를 공식적으로 기록·전달하는 문서. ‘정해진 형식’과 ‘정확한 표현’이 생명입니다.",ORANGE,PALE_ORG)
rules=[("정확성","사실 근거 · 오탈자·숫자 오류 없이"),("간결성","한 문장 한 뜻 · 짧게"),("명확성","애매한 표현 금지"),("성실·예의","정중한 표현 · 상대 입장"),("형식 준수","두문-본문-결문 · 항목 기호")]
for i,(t1,t2) in enumerate(rules):
    y=Inches(3.55)+Inches(0.6)*i
    rect(s,M,y,Inches(6.0),Inches(0.5),LIGHT,shape=MSO_SHAPE.ROUNDED_RECTANGLE,round_=0.1)
    rect(s,M+Inches(0.15),y+Inches(0.1),Inches(0.3),Inches(0.3),ORANGE,shape=MSO_SHAPE.OVAL)
    text(s,M+Inches(0.15),y+Inches(0.1),Inches(0.3),Inches(0.3),[[(str(i+1),11,WHITE,True)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    text(s,M+Inches(0.65),y,Inches(1.5),Inches(0.5),[[(t1,13.5,NAVY,True)]],anchor=MSO_ANCHOR.MIDDLE)
    text(s,M+Inches(2.1),y,Inches(3.8),Inches(0.5),[[(t2,12,GRAY,False)]],anchor=MSO_ANCHOR.MIDDLE)
rx=Inches(7.1)
rect(s,rx,Inches(1.95),Inches(5.6),Inches(4.05),NAVY,shape=MSO_SHAPE.ROUNDED_RECTANGLE,round_=0.04)
text(s,rx+Inches(0.45),Inches(2.25),Inches(5),Inches(0.4),[[("항목 구분 순서",16,GOLD,True)]])
text(s,rx+Inches(0.45),Inches(2.95),Inches(4.8),Inches(1.5),
     [[("1.  →  가.  →  1)  →  가)",24,WHITE,True)],[("→  (1)  →  (가)  →  ①  →  ㉮",24,WHITE,True)]],line_spacing=1.5)
text(s,rx+Inches(0.45),Inches(4.9),Inches(4.8),Inches(0.9),[[("둘째 항목부터는 상위 항목 위치에서 1자(2타)씩 들여쓰기",13,SUBSKY,False)]],line_spacing=1.2)
footer(s,"형식이 신뢰를 만든다",ORANGE)

# ====================================================================
# 21. 공문서 양식 목업
# ====================================================================
s=slide(); bg(s)
header(s,"PART 2 · 기본 이론","공문서 양식 한눈에 보기",ORANGE)
PGB=RGBColor(0xB4,0xBC,0xC8)
px,pw=Inches(0.95),Inches(4.6); py,ph=Inches(1.95),Inches(4.55)
rect(s,px,py,pw,ph,WHITE,line=PGB,line_w=1.3,shadow=True)
ix=px+Inches(0.3); iw=pw-Inches(0.6)
def bdg(n,cy):
    d=Inches(0.34)
    rect(s,px-Inches(0.5),cy,d,d,ORANGE,shape=MSO_SHAPE.OVAL)
    text(s,px-Inches(0.5),cy,d,d,[[(str(n),11,WHITE,True)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
text(s,px,py+Inches(0.15),pw,Inches(0.4),[[("한전원자력연료(주)",15,NAVY,True)]],align=PP_ALIGN.CENTER); bdg(1,py+Inches(0.16))
rect(s,ix,py+Inches(0.62),iw,Pt(2),NAVY); rect(s,ix,py+Inches(0.68),iw,Pt(1),NAVY)
text(s,ix,py+Inches(0.85),iw,Inches(0.3),[[("수신  ○○○○장  (○○과장)",10.5,INK,False)]]); bdg(2,py+Inches(0.83))
text(s,ix,py+Inches(1.25),iw,Inches(0.3),[[("제목  ",11,ORANGE,True),("신입직원 입문교육 결과 보고",11,NAVY,True)]]); bdg(3,py+Inches(1.23))
text(s,ix,py+Inches(1.68),iw,Inches(0.3),[[("1. 신입직원 입문교육 결과를 다음과 같이 보고합니다.",9.3,INK,False)]])
text(s,ix+Inches(0.2),py+Inches(2.0),iw,Inches(0.3),[[("가. 교육 개요 / 나. 주요 결과",8.8,GRAY,False)]]); bdg(4,py+Inches(1.7))
text(s,ix,py+Inches(2.45),iw,Inches(0.3),[[("붙임  1. 결과보고서 1부.   ",9.5,INK,False),("끝.",9.5,NAVY,True)]]); bdg(5,py+Inches(2.43))
text(s,px,py+Inches(2.9),pw,Inches(0.4),[[("한전원자력연료 사장",13,NAVY,True),("  (직인)",8.5,GRAY,False)]],align=PP_ALIGN.CENTER); bdg(6,py+Inches(2.95))
rect(s,ix,py+Inches(3.4),iw,Inches(0.4),LIGHT,line=PGB,line_w=1,shape=MSO_SHAPE.ROUNDED_RECTANGLE,round_=0.1)
text(s,ix,py+Inches(3.4),iw,Inches(0.4),[[("기안 ○○○   검토 ○○○   결재 ○○○",9,GRAY,False)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE); bdg(7,py+Inches(3.45))
rect(s,ix,py+Inches(3.95),iw,Pt(1.2),PGB)
text(s,px,py+Inches(4.05),pw,Inches(0.4),[[("대전 유성구 ○○로 000  ·  ☎ 042-000-0000  ·  www.knfc.co.kr",7.8,GRAY,False)]],align=PP_ALIGN.CENTER); bdg(8,py+Inches(4.05))
# legend right
lx=Inches(6.2)
leg=["발신 기관명 (로고)","수신자 (받는 사람)","제목 (한 줄 요약)","본문 (1. 가. 항목 기호)","붙임 + “끝.” 표시","발신 명의·직인","결재란 (기안·검토·결재)","기관 주소·연락처"]
text(s,lx,Inches(1.95),Inches(6),Inches(0.4),[[("①~⑧ 구성 요소",16,ORANGE,True)]])
for i,t1 in enumerate(leg):
    y=Inches(2.5)+Inches(0.48)*i
    rect(s,lx,y,Inches(0.34),Inches(0.34),ORANGE,shape=MSO_SHAPE.OVAL)
    text(s,lx,y,Inches(0.34),Inches(0.34),[[(str(i+1),11,WHITE,True)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    text(s,lx+Inches(0.5),y,Inches(5.5),Inches(0.36),[[(t1,14,INK,False)]],anchor=MSO_ANCHOR.MIDDLE)
footer(s,"정해진 자리에, 정해진 내용을",ORANGE)

# ====================================================================
# 22. 공문서 3단 구조 + 표기법
# ====================================================================
s=slide(); bg(s)
header(s,"PART 2 · 기본 이론","공문서 3단 구조 — 두문·본문·결문",ORANGE)
def zone(y,h,accent,pale,kr,who,items):
    rect(s,M,y,Inches(8.4),h,pale,shape=MSO_SHAPE.ROUNDED_RECTANGLE,round_=0.04)
    rect(s,M,y,Inches(1.6),h,accent,shape=MSO_SHAPE.ROUNDED_RECTANGLE,round_=0.06)
    text(s,M,y+h/2-Inches(0.38),Inches(1.6),Inches(0.5),[[(kr,22,WHITE,True)]],align=PP_ALIGN.CENTER)
    text(s,M,y+h/2+Inches(0.08),Inches(1.6),Inches(0.3),[[(who,11,WHITE,False)]],align=PP_ALIGN.CENTER)
    tx=M+Inches(1.85)
    for j,(k,v) in enumerate(items):
        yy=y+Inches(0.12)+(h-Inches(0.24))/len(items)*j
        text(s,tx,yy,Inches(1.6),Inches(0.4),[[(k,13.5,accent,True)]],anchor=MSO_ANCHOR.MIDDLE)
        text(s,tx+Inches(1.7),yy,Inches(4.6),Inches(0.4),[[(v,13,INK,False)]],anchor=MSO_ANCHOR.MIDDLE)
zone(Inches(1.95),Inches(0.95),BLUE,PALE_BLUE,"두문","누구에게",[("기관명","발신 기관"),("수신·경유","받는 사람 / 거쳐가는 부서")])
zone(Inches(3.0),Inches(1.45),TEAL,PALE_TEAL,"본문","무엇을",[("제목","내용을 한 줄로 (명사형)"),("내용","1.→가.→1) 항목 기호 순서"),("붙임","첨부 + 본문 끝 “끝.”")])
zone(Inches(4.55),Inches(1.3),ORANGE,PALE_ORG,"결문","누가·언제",[("발신 명의","○○기관장 + 직인"),("결재·시행","기안·검토·결재 / 시행일")])
# cheat
text(s,M,Inches(6.0),Inches(6),Inches(0.35),[[("✍️ 자주 틀리는 표기법",14,NAVY,True)]])
ch=[("날짜","2026. 7. 8."),("금액","금113,560원"),("시간","14:00"),("끝 표시","본문 끝 “끝.”")]
for i,(t1,ex) in enumerate(ch):
    x=M+(Inches(2.95)+Inches(0.12))*i
    rect(s,x,Inches(6.4),Inches(2.95),Inches(0.55),PALE_ORG,shape=MSO_SHAPE.ROUNDED_RECTANGLE,round_=0.2)
    text(s,x+Inches(0.25),Inches(6.4),Inches(1.1),Inches(0.55),[[(t1,12,ORANGE,True)]],anchor=MSO_ANCHOR.MIDDLE)
    text(s,x+Inches(1.25),Inches(6.4),Inches(1.6),Inches(0.55),[[(ex,12.5,NAVY,True)]],anchor=MSO_ANCHOR.MIDDLE)
footer(s,"구조를 알면, 공문서가 쉬워진다",ORANGE)

# ====================================================================
# 23. 원페이지 보고서
# ====================================================================
s=slide(); bg(s)
header(s,"PART 2 · 기본 이론","원페이지 보고서 — 한 장으로",ORANGE)
def_box(s,M,Inches(1.95),PW-2*M,Inches(1.0),"왜 한 장인가",
        "바쁜 의사결정자는 ‘한 장’을 원합니다. 길이가 아니라 ‘구성’이 실력 — 핵심만 구조화하면 한 장으로 충분.",ORANGE,PALE_ORG)
struct=[("제목","한눈에 주제 파악"),("개요·목적","왜 이 보고를 하나"),("현황·문제","사실·데이터 중심"),
        ("대안·방안","1~3개 선택지 비교"),("결론·건의","무엇을 승인받을지"),("향후 일정","언제·누가·무엇을")]
w=Inches(3.85)
for i,(t1,t2) in enumerate(struct):
    r=i//3; c=i%3
    x=M+(w+Inches(0.3))*c; y=Inches(3.25)+Inches(1.4)*r
    rect(s,x,y,w,Inches(1.2),WHITE,line=LINE,line_w=1,shape=MSO_SHAPE.ROUNDED_RECTANGLE,round_=0.05,shadow=True)
    rect(s,x,y,Inches(0.14),Inches(1.2),ORANGE,shape=MSO_SHAPE.ROUNDED_RECTANGLE,round_=0.5)
    text(s,x+Inches(0.4),y+Inches(0.2),w-Inches(0.6),Inches(0.45),[[(str(i+1)+". "+t1,17,ORANGE,True)]])
    text(s,x+Inches(0.4),y+Inches(0.68),w-Inches(0.6),Inches(0.4),[[(t2,13,GRAY,False)]])
footer(s,"길게 쓰지 말고, 구조로 압축하라",ORANGE)

# ====================================================================
# 24. AI 보고서 4단계
# ====================================================================
s=slide(); bg(s)
header(s,"PART 2 · AI 활용","AI로 보고서 빠르게 — 4단계",ORANGE)
text(s,M,Inches(1.95),PW-2*M,Inches(0.4),[[("AI는 ‘초안·구조화·다듬기’를 돕는다. ",15,INK,False),("‘판단과 책임’은 나의 몫.",15,ORANGE,True)]])
flow=[("기획","목적·대상·핵심을 한 줄로 정리","AI에게 줄 ‘재료’"),
      ("초안","원페이지 구조로 초안 요청","개요·현황·방안·건의 순"),
      ("검증","사실·수치·근거를 직접 확인","지어낸 내용 점검 (가장 중요!)"),
      ("다듬기","톤·길이·표현을 AI로 마무리","“간결하게/두괄식/공문체”")]
w=Inches(2.95)
for i,(t1,t2,t3) in enumerate(flow):
    x=M+(w+Inches(0.13))*i
    rect(s,x,Inches(2.6),w,Inches(3.3),WHITE,line=LINE,line_w=1,shape=MSO_SHAPE.ROUNDED_RECTANGLE,round_=0.05,shadow=True)
    rect(s,x,Inches(2.6),w,Inches(0.9),ORANGE,shape=MSO_SHAPE.ROUND_2_SAME_RECTANGLE,round_=0.1)
    text(s,x,Inches(2.7),w,Inches(0.7),[[(str(i+1),34,WHITE,True)]],align=PP_ALIGN.CENTER)
    text(s,x,Inches(3.65),w,Inches(0.5),[[(t1,20,NAVY,True)]],align=PP_ALIGN.CENTER)
    text(s,x+Inches(0.3),Inches(4.25),w-Inches(0.6),Inches(1.0),[[(t2,14,INK,False)]],align=PP_ALIGN.CENTER,line_spacing=1.2)
    rect(s,x+Inches(0.25),Inches(5.2),w-Inches(0.5),Inches(0.55),LIGHT,shape=MSO_SHAPE.ROUNDED_RECTANGLE,round_=0.1)
    text(s,x+Inches(0.3),Inches(5.2),w-Inches(0.6),Inches(0.55),[[(t3,11,GRAY,False)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE,line_spacing=1.0)
    if i<3:
        text(s,x+w-Inches(0.05),Inches(3.9),Inches(0.3),Inches(0.5),[[("›",26,ORANGE,True)]])
footer(s,"AI는 거들 뿐 — 판단은 사람이",ORANGE)

# ====================================================================
# 25. 전지+매직 실습
# ====================================================================
s=slide(); bg(s)
header(s,"PART 2 · 팀 실습","전지·매직 보고서 작성",ORANGE)
rect(s,M,Inches(1.95),Inches(6.0),Inches(1.1),PALE_ORG,shape=MSO_SHAPE.ROUNDED_RECTANGLE,round_=0.05)
rect(s,M,Inches(1.95),Inches(0.16),Inches(1.1),ORANGE,shape=MSO_SHAPE.ROUNDED_RECTANGLE,round_=0.5)
text(s,M+Inches(0.4),Inches(2.12),Inches(5),Inches(0.4),[[("🎯 실습 목표",15,ORANGE,True)]])
text(s,M+Inches(0.4),Inches(2.5),Inches(5.4),Inches(0.5),[[("AI 초안 → 조별 ",13.5,INK,False),("전지에 매직으로 원페이지 보고서",13.5,ORANGE,True),(" 완성·발표",13.5,INK,False)]],line_spacing=1.15)
steps=[("주제 선정·기획","목적·대상·핵심 합의 (5분)"),("AI로 초안 (모바일)","구조 초안 → 사실 검증 (10분)"),
       ("전지에 옮겨 구성","원페이지 골격대로 작성 (20분)"),("조별 발표·피드백","1분 두괄식 발표 (5분)")]
for i,(t1,t2) in enumerate(steps):
    y=Inches(3.4)+Inches(0.72)*i
    rect(s,M,y,Inches(6.0),Inches(0.6),LIGHT,shape=MSO_SHAPE.ROUNDED_RECTANGLE,round_=0.06)
    rect(s,M+Inches(0.15),y+Inches(0.1),Inches(0.4),Inches(0.4),ORANGE,shape=MSO_SHAPE.OVAL)
    text(s,M+Inches(0.15),y+Inches(0.1),Inches(0.4),Inches(0.4),[[(str(i+1),13,WHITE,True)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    text(s,M+Inches(0.75),y,Inches(2.4),Inches(0.6),[[(t1,14,NAVY,True)]],anchor=MSO_ANCHOR.MIDDLE)
    text(s,M+Inches(3.1),y,Inches(2.8),Inches(0.6),[[(t2,11.5,GRAY,False)]],anchor=MSO_ANCHOR.MIDDLE)
rx=Inches(7.1)
rect(s,rx,Inches(1.95),Inches(5.6),Inches(4.05),NAVY,shape=MSO_SHAPE.ROUNDED_RECTANGLE,round_=0.04)
text(s,rx+Inches(0.45),Inches(2.25),Inches(5),Inches(0.4),[[("🏆 평가 기준 (각 5점)",16,GOLD,True)]])
for i,(t1,t2) in enumerate([("구성","원페이지 골격을 갖췄는가"),("두괄식","결론·핵심이 먼저 보이는가"),("사실성","사실/의견 구분·근거가 있나"),("전달력","한눈에 이해되는 구성인가")]):
    y=Inches(2.95)+Inches(0.72)*i
    text(s,rx+Inches(0.5),y,Inches(1.8),Inches(0.5),[[("· "+t1,15,WHITE,True)]],anchor=MSO_ANCHOR.MIDDLE)
    text(s,rx+Inches(2.0),y,Inches(3.4),Inches(0.5),[[(t2,12.5,SUBSKY,False)]],anchor=MSO_ANCHOR.MIDDLE)
footer(s,"손으로 쓰면, 머리에 남는다",ORANGE)

# ====================================================================
# 26. 미션카드
# ====================================================================
s=slide(); bg(s)
header(s,"PART 2 · 팀 실습","보고서 미션카드 (택 1)",ORANGE)
text(s,M,Inches(1.9),PW-2*M,Inches(0.4),[[("조별로 1개를 골라 ",14,INK,False),("자유롭게 토론하며",14,ORANGE,True),(" 원페이지 보고서로 완성하세요.",14,INK,False)]])
cards=[("📋","A · 행사기획","신입사원 워크숍 개최 계획"),("🛠️","B · 업무개선","반복 업무 자동화 제안"),("📊","C · 현황보고","교육 만족도 조사 결과"),
       ("⚠️","D · 상황보고","업무 지연 중간보고"),("💡","E · 아이디어톤","사내 문화·복지 개선"),("📣","F · 홍보기획","MZ 겨냥 회사 SNS 홍보")]
w=Inches(3.85)
for i,(ic,tag,t1) in enumerate(cards):
    r=i//3; c=i%3
    x=M+(w+Inches(0.3))*c; y=Inches(2.5)+Inches(1.7)*r
    rect(s,x,y,w,Inches(1.5),WHITE,line=LINE,line_w=1,shape=MSO_SHAPE.ROUNDED_RECTANGLE,round_=0.05,shadow=True)
    rect(s,x,y,w,Inches(0.5),PALE_ORG,shape=MSO_SHAPE.ROUND_2_SAME_RECTANGLE,round_=0.2)
    text(s,x+Inches(0.3),y,w-Inches(0.6),Inches(0.5),[[("MISSION "+tag,13,ORANGE,True)]],anchor=MSO_ANCHOR.MIDDLE)
    text(s,x+Inches(0.3),y+Inches(0.6),Inches(0.7),Inches(0.7),[[(ic,26,INK,False)]],anchor=MSO_ANCHOR.MIDDLE)
    text(s,x+Inches(1.15),y+Inches(0.6),w-Inches(1.4),Inches(0.7),[[(t1,15.5,NAVY,True)]],anchor=MSO_ANCHOR.MIDDLE,line_spacing=1.1)
footer(s,"어떤 일이든, 보고는 한 장이면 된다",ORANGE)

# ====================================================================
# 27. WRAP-UP
# ====================================================================
s=slide(); bg(s)
header(s,"WRAP-UP · 핵심 정리","오늘 챙겨갈 5가지",BLUE,GOLD)
take=[("AI는 ‘도구’ — 일을 더 빠르게 돕는다",TEAL),("좋은 질문이 좋은 결과를 만든다 (R·O·C·F·E·C)",TEAL),
      ("결과물은 ‘초안’ — 사실·수치는 내가 검증",ORANGE),("보고는 두괄식 — 결론부터, 사실/의견 구분",ORANGE),
      ("보안 최우선 — 민감·대외비는 공개형 AI 금지",RED)]
for i,(t1,col) in enumerate(take):
    y=Inches(1.95)+Inches(0.92)*i
    rect(s,M,y,PW-2*M,Inches(0.78),WHITE,line=LINE,line_w=1,shape=MSO_SHAPE.ROUNDED_RECTANGLE,round_=0.06,shadow=True)
    rect(s,M,y,Inches(0.16),Inches(0.78),col,shape=MSO_SHAPE.ROUNDED_RECTANGLE,round_=0.5)
    rect(s,M+Inches(0.4),y+Inches(0.19),Inches(0.4),Inches(0.4),col,shape=MSO_SHAPE.OVAL)
    text(s,M+Inches(0.4),y+Inches(0.19),Inches(0.4),Inches(0.4),[[(str(i+1),15,WHITE,True)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    text(s,M+Inches(1.1),y,PW-2*M-Inches(1.4),Inches(0.78),[[(t1,16.5,INK,True)]],anchor=MSO_ANCHOR.MIDDLE)
footer(s,"오늘 배운 한 가지를, 내일 업무에서",GOLD)

# ====================================================================
# 28. THANK YOU
# ====================================================================
s=slide(); bg(s,NAVY)
rect(s,0,0,PW,Inches(0.24),TEAL); rect(s,0,Inches(0.24),PW,Inches(0.08),ORANGE)
rect(s,Inches(9.0),Inches(1.2),Inches(4.5),Inches(4.5),RGBColor(0x16,0x3A,0x66),shape=MSO_SHAPE.OVAL)
text(s,M,Inches(2.2),Inches(8),Inches(0.5),[[("ACTION PLAN",16,GOLD,True)]])
text(s,M,Inches(2.7),Inches(9),Inches(0.8),[[("내일부터 이렇게 써보겠습니다",30,WHITE,True)]])
rect(s,M,Inches(3.6),Inches(1.2),Inches(0.07),ORANGE)
for i,(t1,t2,col) in enumerate([("AI 활용 1가지","예: 회의록 요약을 노트북LM으로",TEAL),("보고 습관 1가지","예: 모든 보고는 결론부터",ORANGE),("이번 주 실천","예: 주간보고를 AI 초안→검증으로",GOLD)]):
    x=M+(Inches(3.95)+Inches(0.1))*i
    rect(s,x,Inches(3.9),Inches(3.95),Inches(1.3),CARDNAVY,shape=MSO_SHAPE.ROUNDED_RECTANGLE,round_=0.05)
    rect(s,x+Inches(0.3),Inches(4.1),Inches(0.5),Inches(0.5),col,shape=MSO_SHAPE.OVAL)
    text(s,x+Inches(0.3),Inches(4.1),Inches(0.5),Inches(0.5),[[(str(i+1),16,WHITE,True)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    text(s,x+Inches(0.95),Inches(4.12),Inches(2.8),Inches(0.4),[[(t1,15,WHITE,True)]])
    text(s,x+Inches(0.3),Inches(4.7),Inches(3.4),Inches(0.45),[[(t2,11,SUBSKY,False)]],line_spacing=1.0)
text(s,M,Inches(5.8),PW-2*M,Inches(0.8),[[("Thank You",38,WHITE,True)]],align=PP_ALIGN.CENTER)
text(s,M,Inches(6.6),PW-2*M,Inches(0.4),[[("한전원자력연료 · 2026 신입직원 입문교육",13,SUBSKY,False)]],align=PP_ALIGN.CENTER)

OUT="/home/user/yja-3fire/knf-onboarding-ai-training/한전원자력연료_신입직원_입문교육_강사교안_16x9.pptx"
prs.save(OUT)
print("SAVED", OUT, "slides:", len(prs.slides._sldIdLst))
