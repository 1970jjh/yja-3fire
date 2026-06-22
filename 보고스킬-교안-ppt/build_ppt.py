#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""보고 스킬 실습 교안 — PowerPoint(.pptx) 빌더.

워크북 PDF와 동일한 디자인 시스템(컬러·위계)을 16:9 슬라이드로 옮긴 발표용 교안.
편집 가능한 .pptx 를 생성한다.  실행:  python3 build_ppt.py
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---------- 디자인 토큰 ----------
INK        = RGBColor(0x16,0x22,0x3A)
PRIMARY    = RGBColor(0x2B,0x4C,0x8C)
PRIMARY_DK = RGBColor(0x1E,0x35,0x65)
PRIMARY_SF = RGBColor(0xEA,0xF0,0xFB)
ACCENT     = RGBColor(0xF2,0x6A,0x4B)
ACCENT_SF  = RGBColor(0xFD,0xEB,0xE5)
TEAL       = RGBColor(0x1B,0x9C,0x85)
TEAL_SF    = RGBColor(0xE4,0xF4,0xEF)
AMBER      = RGBColor(0xE8,0x93,0x0C)
AMBER_SF   = RGBColor(0xFD,0xF1,0xDC)
LINE       = RGBColor(0xDC,0xE3,0xEE)
MUTED      = RGBColor(0x6A,0x74,0x88)
TEXT       = RGBColor(0x2A,0x31,0x42)
PAPER      = RGBColor(0xF5,0xF8,0xFC)
WHITE      = RGBColor(0xFF,0xFF,0xFF)
FONT = "맑은 고딕"

SW, SH = Inches(13.333), Inches(7.5)
MX = Inches(0.7)            # 좌우 여백
CW = Inches(13.333-1.4)     # 콘텐츠 폭

prs = Presentation()
prs.slide_width = SW
prs.slide_height = SH
BLANK = prs.slide_layouts[6]


def slide(bg=WHITE):
    s = prs.slides.add_slide(BLANK)
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0,0, SW, SH)
    r.fill.solid(); r.fill.fore_color.rgb = bg
    r.line.fill.background()
    r.shadow.inherit = False
    return s


def rect(s, l,t,w,h, fill=None, line=None, line_w=1.0, rounded=False):
    shp = s.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, l,t,w,h)
    if rounded:
        try: shp.adjustments[0] = 0.06
        except Exception: pass
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line; shp.line.width = Pt(line_w)
    shp.shadow.inherit = False
    return shp


def text(s, l,t,w,h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         space=4, line_sp=1.0, wrap=True):
    """runs: list of paragraphs; each paragraph = list of (txt,size,color,bold)."""
    tb = s.shapes.add_textbox(l,t,w,h)
    tf = tb.text_frame; tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    for m in ("left","right","top","bottom"):
        setattr(tf, "margin_"+m, Emu(0))
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(space); p.space_before = Pt(0)
        p.line_spacing = line_sp
        if isinstance(para, tuple): para=[para]
        for (txt,size,color,bold) in para:
            r = p.add_run(); r.text = txt
            r.font.size = Pt(size); r.font.bold = bold
            r.font.name = FONT; r.font.color.rgb = color
    return tb


def page_no(s, n):
    text(s, SW-Inches(1.0), SH-Inches(0.5), Inches(0.6), Inches(0.3),
         [(str(n), 11, PRIMARY, True)], align=PP_ALIGN.RIGHT)
    rect(s, MX, SH-Inches(0.5), Inches(2.6), Pt(0))  # spacer noop

PAGE = [0]
def footer(s, title="보고 스킬 실습 교안"):
    PAGE[0]+=1
    text(s, MX, SH-Inches(0.5), Inches(6), Inches(0.3),
         [(title, 8.5, RGBColor(0x9A,0xA4,0xB6), False)], anchor=MSO_ANCHOR.MIDDLE)
    text(s, SW-Inches(1.1), SH-Inches(0.5), Inches(0.7), Inches(0.3),
         [(str(PAGE[0]), 10.5, PRIMARY, True)], align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)


def header(s, label, title):
    """콘텐츠 슬라이드 상단 헤더. 반환: 콘텐츠 시작 y."""
    rect(s, 0,0, SW, Inches(1.18), fill=INK)
    rect(s, 0, Inches(1.18), SW, Pt(3), fill=ACCENT)
    text(s, MX, Inches(0.2), Inches(8), Inches(0.3),
         [(label, 11, RGBColor(0xFF,0xC9,0xBC), True)])
    text(s, MX, Inches(0.46), CW, Inches(0.6),
         [(title, 23, WHITE, True)])
    return Inches(1.5)


def block_title(s, y, txt, color=INK):
    text(s, MX, y, CW, Inches(0.4), [(txt, 15, color, True)])
    return y + Inches(0.5)


# ===================================================================
#  슬라이드 구성
# ===================================================================

# ---- 1. 타이틀 ----
s = slide(PRIMARY_DK)
rect(s, 0,0, SW, SH, fill=PRIMARY_DK)
rect(s, 0, Inches(2.5), SW, Pt(4), fill=None)
rect(s, Inches(0.9), Inches(2.0), Inches(0.9), Pt(6), fill=ACCENT)
text(s, Inches(0.9), Inches(1.15), Inches(11), Inches(0.5),
     [("REPORTING SKILLS", 15, RGBColor(0xFF,0xC9,0xBC), True)])
text(s, Inches(0.9), Inches(2.3), Inches(11.5), Inches(2.0),
     [("보고 스킬 실습 교안", 46, WHITE, True)])
text(s, Inches(0.9), Inches(4.2), Inches(11), Inches(1.2),
     [("애매한 상사의 말을 번역하고, 나를 신뢰하게 만드는", 18, RGBColor(0xC9,0xD6,0xEE), False),
      ("신입사원의 보고 기술 — 실습 중심 발표 교안", 18, RGBColor(0xC9,0xD6,0xEE), False)],
     line_sp=1.3)
for i,(c) in enumerate(["대상 · 신입~3년차","소요 · 90분","구성 · 5개 모듈 · 실습 포함"]):
    bx = rect(s, Inches(0.9)+Inches(3.05)*i, Inches(6.2), Inches(2.85), Inches(0.5),
              fill=None, line=RGBColor(0x5A,0x70,0x9A), line_w=1.2, rounded=True)
    text(s, Inches(0.9)+Inches(3.05)*i, Inches(6.3), Inches(2.85), Inches(0.32),
         [(c, 11, RGBColor(0xEA,0xF0,0xFB), True)], align=PP_ALIGN.CENTER)

# ---- 2. 목차 ----
s = slide()
y = header(s, "AGENDA", "오늘 배울 5가지")
items = [
    ("1","상사의 말 번역하기 — 질문의 기술","모르면 그 자리에서 묻는다"),
    ("2","구두 보고 & 과제 프레이밍","타이밍과 기대 수준 맞추기"),
    ("3","이메일 보고 & 보고서 보고","맥락과 의견을 담는다"),
    ("4","보고의 자세 — 메신저 vs 플레이어","대안을 들고 간다"),
    ("5","상사 유형 파악 & 맞춤 소통","유형을 알면 합이 맞는다"),
]
cy = y
for num,t1,t2 in items:
    rect(s, MX, cy, CW, Inches(0.92), fill=PAPER, rounded=True)
    rect(s, MX+Inches(0.16), cy+Inches(0.16), Inches(0.6), Inches(0.6), fill=PRIMARY, rounded=True)
    text(s, MX+Inches(0.16), cy+Inches(0.2), Inches(0.6), Inches(0.5),
         [(num, 20, WHITE, True)], align=PP_ALIGN.CENTER)
    text(s, MX+Inches(1.0), cy+Inches(0.13), Inches(10), Inches(0.4),
         [(t1, 15, INK, True)])
    text(s, MX+Inches(1.0), cy+Inches(0.52), Inches(10), Inches(0.3),
         [(t2, 11, MUTED, False)])
    cy += Inches(1.04)
footer(s)


def divider(num, label, title_lines, desc):
    s = slide()
    rect(s, 0,0, Inches(0.32), SH, fill=ACCENT)
    text(s, Inches(0.9), Inches(1.7), Inches(6), Inches(1.6),
         [(num, 80, PRIMARY_SF, True)])
    text(s, Inches(0.95), Inches(3.25), Inches(10), Inches(0.4),
         [(label, 13, ACCENT, True)])
    text(s, Inches(0.95), Inches(3.65), Inches(11.5), Inches(1.6),
         [[(t, 32, INK, True)] for t in title_lines], line_sp=1.05)
    text(s, Inches(0.95), Inches(5.5), Inches(10.8), Inches(1.3),
         [(desc, 14, MUTED, False)], line_sp=1.35)
    return s

# ---- 3. Divider 1 ----
divider("01","MODULE 01",["상사의 말 번역하기","— 질문의 기술"],
        "상사는 늘 명확히 지시하지 않는다. 심지어 본인도 일을 정확히 모를 때가 있다. 가장 좋은 번역기는 '질문'이며, 신입에게는 질문할 권리가 있다.")

# ---- 4. M1 핵심개념 ----
s = slide()
y = header(s, "MODULE 01 · 핵심개념", "질문이 신뢰를 만든다")
rect(s, MX, y, CW, Inches(1.5), fill=PRIMARY_SF, rounded=True)
rect(s, MX, y, Pt(5), Inches(1.5), fill=PRIMARY)
text(s, MX+Inches(0.3), y+Inches(0.22), CW-Inches(0.6), Inches(1.1),
     [[("지시받은 ",13,TEXT,False),("그 자리에서",13,PRIMARY_DK,True),(" 모르는 부분을 묻지 않는 것이 신입의 치명적 실수다.",13,TEXT,False)],
      [("질문 없이 자리로 돌아가면 상사는 오히려 ",13,TEXT,False),("불안",13,PRIMARY_DK,True),("해진다.",13,TEXT,False)]],
     line_sp=1.25)
y2 = y+Inches(1.8)
col = (CW-Inches(0.3))/2
for i,(tt,dd,c,csf) in enumerate([
    ("① 멋대로 처리할 확률 ↓","모르는 부분을 임의로 처리하다 생기는 사고를 막는다",TEAL,TEAL_SF),
    ("② 적극적이라는 인상 ↑","맡은 일을 능동적으로 대한다는 신뢰를 준다",ACCENT,ACCENT_SF)]):
    x = MX + (col+Inches(0.3))*i
    rect(s, x, y2, col, Inches(1.7), fill=csf, rounded=True)
    text(s, x+Inches(0.3), y2+Inches(0.28), col-Inches(0.6), Inches(0.5),
         [(tt, 16, c, True)])
    text(s, x+Inches(0.3), y2+Inches(0.85), col-Inches(0.6), Inches(0.7),
         [(dd, 13, TEXT, False)], line_sp=1.2)
footer(s)

# ---- 5. M1 사례 (대화) ----
def dialogue_slide(label, title, lines, note=None):
    s = slide()
    y = header(s, label, title)
    cy = y
    for who, say, junior, inner in lines:
        h = Inches(0.78) if not inner else Inches(1.05)
        rect(s, MX+Inches(1.55), cy, CW-Inches(1.55), h, fill=PAPER, rounded=True)
        rect(s, MX, cy, Inches(1.4), Inches(0.5), fill=(TEAL if junior else PRIMARY), rounded=True)
        text(s, MX, cy+Inches(0.08), Inches(1.4), Inches(0.36),
             [(who, 12, WHITE, True)], align=PP_ALIGN.CENTER)
        runs = [(say, 13.5, TEXT, False)]
        text(s, MX+Inches(1.78), cy+Inches(0.14), CW-Inches(2.0), Inches(0.5),
             [runs])
        if inner:
            text(s, MX+Inches(1.78), cy+Inches(0.6), CW-Inches(2.0), Inches(0.4),
                 [[("(속마음) ",11,MUTED,True),(inner,11,MUTED,False)]])
        cy += h + Inches(0.18)
    if note:
        text(s, MX, cy+Inches(0.05), CW, Inches(0.8), [(note,13,MUTED,False)], line_sp=1.2)
    footer(s)
    return s

dialogue_slide("MODULE 01 · 사례", "정 사원의 원형탈모 사건",
    [("박 차장","정 사원, 기획팀에서 '필수 교육 시간' 체크 중인데 이것 좀 처리해줘요.",False,None),
     ("정 사원","넵!",True,"필수 교육 시간이 뭐지? 일단 알겠다고 하고 나중에 알아보자…")],
    note="→ 묻지 않고 '봉사활동 시간'으로 짐작해 야근까지 했지만, 다음 날 \"내가 언제 봉사활동을 말했어요?\"라는 핀잔만. 한 번 물었다면 없었을 고생이다.")

# ---- 6. M1 실습 (before/after) ----
def ba_slide(label, title, goal, before, after_label, after_lines, hint=None):
    s = slide()
    y = header(s, label, title)
    rect(s, MX, y, CW, Inches(0.6), fill=PAPER, rounded=True)
    text(s, MX+Inches(0.25), y+Inches(0.13), CW-Inches(0.5), Inches(0.4),
         [[("실습 목표.  ",12,PRIMARY_DK,True),(goal,12,TEXT,False)]])
    y2 = y+Inches(0.85)
    col = (CW-Inches(0.3))/2
    hh = Inches(3.0)
    # before
    rect(s, MX, y2, col, hh, fill=ACCENT_SF, rounded=True)
    text(s, MX+Inches(0.28), y2+Inches(0.2), col-Inches(0.5), Inches(0.4),
         [("BEFORE — 이렇게 하면 고생", 12, ACCENT, True)])
    text(s, MX+Inches(0.28), y2+Inches(0.75), col-Inches(0.56), hh-Inches(1.0),
         [[(before,13.5,TEXT,False)]], line_sp=1.3)
    # after
    x2 = MX+col+Inches(0.3)
    rect(s, x2, y2, col, hh, fill=TEAL_SF, rounded=True)
    text(s, x2+Inches(0.28), y2+Inches(0.2), col-Inches(0.5), Inches(0.4),
         [(after_label, 12, TEAL, True)])
    text(s, x2+Inches(0.28), y2+Inches(0.75), col-Inches(0.56), hh-Inches(1.0),
         [[(a,13.5,TEXT,False)] for a in after_lines], line_sp=1.3)
    if hint:
        ry = y2+hh+Inches(0.18)
        rect(s, MX, ry, CW, Inches(0.85), fill=AMBER_SF, rounded=True)
        rect(s, MX, ry, Pt(5), Inches(0.85), fill=AMBER)
        text(s, MX+Inches(0.28), ry+Inches(0.13), CW-Inches(0.6), Inches(0.6),
             [[("힌트.  ",12,AMBER,True),(hint,12,TEXT,False)]], line_sp=1.2)
    footer(s)
    return s

ba_slide("MODULE 01 · 실습 1", "\"넵\"을 똑똑한 질문으로 바꾸기",
    "돌아서서 혼자 추측하는 대답을, 맥락을 파악하는 질문으로 바꾼다.",
    "\"넵.\"  (돌아서서 혼자 추측 → 엉뚱한 결과물)",
    "AFTER — 모범 질문",
    ["\"네. 그런데 '필수 교육 시간'을 조금 더 설명해 주실 수 있을까요?",
     "그리고 이걸 체크하는 목적은 무엇인가요?\"",
     "",
     "(센스 버전) \"정확한 맥락 파악을 위해 기획팀에서 받으신 메일을 전달해 주시겠어요?\""],
    hint="① 목적을 묻고  ② '기획팀'이라는 단어를 놓치지 말 것 — 일의 출처로 질문을 넘기면 상사도 편하고 실수도 준다.")

# ---- 7. M1 정리 ----
def summary_slide(label, title, points):
    s = slide()
    y = header(s, label, title)
    rect(s, MX, y, CW, Inches(4.6), fill=INK, rounded=True)
    text(s, MX+Inches(0.45), y+Inches(0.35), CW-Inches(0.9), Inches(0.5),
         [("핵심 정리", 16, ACCENT, True)])
    runs=[]
    for p in points:
        runs.append([("•  ",15,ACCENT,True)]+p)
    text(s, MX+Inches(0.45), y+Inches(1.05), CW-Inches(0.9), Inches(3.3),
         runs, line_sp=1.2, space=12)
    footer(s)
    return s

summary_slide("MODULE 01 · 정리", "질문의 기술 — 핵심 3",
    [[("상사의 지시는 ",15,RGBColor(0xE6,0xEC,0xF6),False),("그 자리에서 질문",15,WHITE,True),("해 해석한다",15,RGBColor(0xE6,0xEC,0xF6),False)],
     [("상사도 모르면 ",15,RGBColor(0xE6,0xEC,0xF6),False),("직접 관련된 부서",15,WHITE,True),("에 질문한다",15,RGBColor(0xE6,0xEC,0xF6),False)],
     [("메일 등으로 ",15,RGBColor(0xE6,0xEC,0xF6),False),("맥락",15,WHITE,True),("을 파악하면 더 좋은 질문을 할 수 있다",15,RGBColor(0xE6,0xEC,0xF6),False)]])

# ---- 8. Divider 2 ----
divider("02","MODULE 02",["구두 보고 &","과제 프레이밍"],
        "의사결정이 빨라진 시대, 보고는 '신속함'이 핵심이다. 구두 보고는 타이밍이 생명, 과제 프레이밍으로 기대 수준을 미리 맞추면 헛수고를 막는다.")

# ---- 9. M2 타이밍 ----
s = slide()
y = header(s, "MODULE 02 · 핵심개념", "구두 보고의 4가지 타이밍")
tiles = [("①","업무를 지시받았을 때"),("②","개괄 조사를 마쳤을 때"),
         ("③","방향·일정을 조율할 때"),("④","문서로 남기기 전에")]
col=(CW-Inches(0.45))/2
for i,(n,t1) in enumerate(tiles):
    r=i//2; c=i%2
    x=MX+(col+Inches(0.45))*c; ty=y+(Inches(1.4)+Inches(0.3))*r
    rect(s,x,ty,col,Inches(1.4),fill=PAPER,rounded=True)
    rect(s,x,ty,Pt(5),Inches(1.4),fill=PRIMARY)
    text(s,x+Inches(0.3),ty+Inches(0.32),Inches(0.8),Inches(0.8),[(n,30,PRIMARY,True)])
    text(s,x+Inches(1.15),ty+Inches(0.46),col-Inches(1.3),Inches(0.6),[(t1,16,INK,True)])
ry=y+Inches(3.55)
rect(s,MX,ry,CW,Inches(0.85),fill=ACCENT_SF,rounded=True)
rect(s,MX,ry,Pt(5),Inches(0.85),fill=ACCENT)
text(s,MX+Inches(0.28),ry+Inches(0.13),CW-Inches(0.6),Inches(0.6),
     [[("주의.  ",12,ACCENT,True),("\"지금 준비 중입니다\"만 반복하면 상사는 속으로 벼른다. 진행 정도와 막힌 지점을 말하는 것이 보고다.",12,TEXT,False)]],line_sp=1.2)
footer(s)

# ---- 10. M2 과제 프레이밍 ----
s = slide()
y = header(s, "MODULE 02 · 핵심개념", "과제 프레이밍 3단계")
steps=[("1  목적 정의","받은 당일 목적·용도를 물어 의도를 파악한다. 관계없는 정보를 자연스레 거를 수 있다."),
       ("2  개괄 조사","빠른 데스크 리서치로 결과물의 '감'을 잡고 기대 수준을 상사와 공유한다."),
       ("3  프레이밍 수정","방향·기간 변화 시 ①상황 전달 ②수정안 제안 ③의견 요청 순으로 보고한다.")]
cy=y
for t1,d in steps:
    rect(s,MX,cy,CW,Inches(1.25),fill=TEAL_SF,rounded=True)
    rect(s,MX,cy,Pt(5),Inches(1.25),fill=TEAL)
    text(s,MX+Inches(0.35),cy+Inches(0.2),Inches(4),Inches(0.5),[(t1,17,TEAL,True)])
    text(s,MX+Inches(0.35),cy+Inches(0.72),CW-Inches(0.7),Inches(0.5),[(d,13,TEXT,False)],line_sp=1.15)
    cy+=Inches(1.42)
footer(s)

# ---- 11. M2 실습 ----
s = slide()
y = header(s, "MODULE 02 · 실습", "프레이밍 수정 보고 만들기")
rect(s, MX, y, CW, Inches(1.0), fill=PAPER, rounded=True)
text(s, MX+Inches(0.25), y+Inches(0.15), CW-Inches(0.5), Inches(0.8),
     [[("상황.  ",12.5,PRIMARY_DK,True),("\"내일모레 오전까지 드리기로 했는데, 같은 시간에 신입 전체 교육이 잡혔다는 공지가 왔다.\"",12.5,TEXT,False)],
      [("→ ①상황 ②제안 ③의견요청 3단 구조로 한 문장 보고를 만들어 보자.",12,MUTED,False)]],line_sp=1.25)
y2=y+Inches(1.25)
rect(s, MX, y2, CW, Inches(1.4), fill=WHITE, line=LINE, line_w=1.4, rounded=True)
text(s, MX+Inches(0.3), y2+Inches(0.25), CW-Inches(0.6), Inches(0.9),
     [("작성란 ✎  ___________________________________________________________", 13, MUTED, False)])
ry=y2+Inches(1.65)
rect(s, MX, ry, CW, Inches(1.2), fill=TEAL_SF, rounded=True)
text(s, MX+Inches(0.3), ry+Inches(0.18), CW-Inches(0.6), Inches(0.4),
     [("모범 예시", 13, TEAL, True)])
text(s, MX+Inches(0.3), ry+Inches(0.6), CW-Inches(0.6), Inches(0.5),
     [[("\"신입 전체 교육이 내일모레 오전에 잡혔습니다(상황). 혹시 ",12.5,TEXT,False),
       ("그다음 날 오전까지",12.5,PRIMARY_DK,True),
       (" 드려도 괜찮을까요(제안)? 일정 조정이 가능한지 여쭙습니다(의견요청).\"",12.5,TEXT,False)]],line_sp=1.2)
footer(s)

# ---- 12. M2 정리 ----
summary_slide("MODULE 02 · 정리", "구두 보고 — 핵심 3",
    [[("구두 보고는 ",15,RGBColor(0xE6,0xEC,0xF6),False),("타이밍",15,WHITE,True),(" — 지시 직후·개괄조사 후·조율 시·문서화 전",15,RGBColor(0xE6,0xEC,0xF6),False)],
     [("과제 프레이밍",15,WHITE,True),("으로 결과물의 기대 수준을 미리 맞춘다",15,RGBColor(0xE6,0xEC,0xF6),False)],
     [("프레임 변경은 ",15,RGBColor(0xE6,0xEC,0xF6),False),("상황 → 제안 → 의견요청",15,WHITE,True),(" 순으로 보고한다",15,RGBColor(0xE6,0xEC,0xF6),False)]])

# ---- 13. Divider 3 ----
divider("03","MODULE 03",["이메일 보고 &","보고서 보고"],
        "이메일은 빠르고 체계적이지만 맥락이 빠지기 쉽다. 보고서는 완결된 콘텐츠인 만큼 목적·두괄식·의견이 핵심이다.")

# ---- 14. M3 이메일 3원칙 ----
s = slide()
y = header(s, "MODULE 03 · 핵심개념", "이메일 보고 3원칙")
rows=[("① 제목에 목적","'Re:'·'안녕하세요'가 아니라 \"경쟁사 현황 데이터 및 동향 보고\"처럼."),
      ("② 서론·본론·결론","서론(경위·목적) · 본론(자료 상세) · 결론(1차 의견·제안)."),
      ("③ 수신·참조 구별","받는 사람=실행자, 참조=관계자. 섞이면 서로 미룬다.")]
cy=y
for t1,d in rows:
    rect(s,MX,cy,CW,Inches(1.25),fill=PRIMARY_SF,rounded=True)
    rect(s,MX,cy,Pt(5),Inches(1.25),fill=PRIMARY)
    text(s,MX+Inches(0.35),cy+Inches(0.2),Inches(6),Inches(0.5),[(t1,17,PRIMARY_DK,True)])
    text(s,MX+Inches(0.35),cy+Inches(0.72),CW-Inches(0.7),Inches(0.5),[(d,13,TEXT,False)],line_sp=1.15)
    cy+=Inches(1.42)
footer(s)

# ---- 15. M3 이메일 before/after ----
ba_slide("MODULE 03 · 실습", "맥락 없는 이메일 살리기",
    "첨부만 덜렁 보내지 말 것. 빠진 3가지를 채워 본문을 완성한다.",
    "제목: Re: 경쟁사\n\n\"말씀하신 자료입니다. 확인 부탁드립니다. (엑셀 첨부)\"",
    "AFTER — 맥락을 담은 보고",
    ["제목: 경쟁사 현황 데이터 및 동향 보고",
     "",
     "① 첨부 설명: 시트1=뉴스 요약 / 시트2=증권사 리포트",
     "② 1차 의견: 주목할 특이점은 B사의 대규모 투자",
     "③ 추가 정보: 미국 신제품 컨퍼런스 기사 동봉"],
    hint="이메일은 맥락 전달이 부족해지기 쉽다 — 첨부 설명·보고자 의견·추가 정보를 함께 담으면 추가 질문이 사라진다.")

# ---- 16. M3 보고서 3원칙 + So What ----
s = slide()
y = header(s, "MODULE 03 · 핵심개념", "보고서 3원칙 & 'So What?'")
col=(CW-Inches(0.4))/2
items=[("① 목적 먼저","왜·어디에 쓰는지 확인 후 시작. 어긋나면 손해가 크다.",PRIMARY,PRIMARY_SF),
       ("② 두괄식","첫 페이지에 핵심·결과. 근거는 뒤에 배치한다.",PRIMARY,PRIMARY_SF),
       ("③ 의견을 담아라","신입에겐 질문할 권리 + '의견 낼 권리'가 있다.",TEAL,TEAL_SF)]
for i,(t1,d,c,sf) in enumerate(items):
    if i<2:
        x=MX+(col+Inches(0.4))*i; ty=y
        w=col
    else:
        x=MX; ty=y+Inches(1.55); w=CW
    rect(s,x,ty,w,Inches(1.4),fill=sf,rounded=True)
    rect(s,x,ty,Pt(5),Inches(1.4),fill=c)
    text(s,x+Inches(0.3),ty+Inches(0.22),w-Inches(0.6),Inches(0.5),[(t1,16,c,True)])
    text(s,x+Inches(0.3),ty+Inches(0.78),w-Inches(0.6),Inches(0.5),[(d,13,TEXT,False)],line_sp=1.15)
ry=y+Inches(3.25)
rect(s,MX,ry,CW,Inches(0.95),fill=ACCENT_SF,rounded=True)
rect(s,MX,ry,Pt(5),Inches(0.95),fill=ACCENT)
text(s,MX+Inches(0.3),ry+Inches(0.16),CW-Inches(0.6),Inches(0.7),
     [[("피해야 할 보고.  ",12.5,ACCENT,True),("점수만 나열 → \"그래서 어떤 안이 적합하지?\"를 되묻게 한다. ",12.5,TEXT,False),("제안까지가 보고자의 임무.",12.5,INK,True)]],line_sp=1.2)
footer(s)

# ---- 17. Divider 4 ----
divider("04","MODULE 04",["보고의 자세","— 메신저 vs 플레이어"],
        "상사는 보고의 '자세'로 부하의 일하는 방식을 유추한다. 실수는 용서해도 소극적 태도엔 미운 정도 안 생긴다. 단순 전달자가 아니라 문제를 푸는 선수가 되자.")

# ---- 18. M4 매트릭스 ----
s = slide()
y = header(s, "MODULE 04 · 핵심개념", "메신저 vs 플레이어")
col=(CW-Inches(0.4))/2
rect(s,MX,y,col,Inches(2.4),fill=ACCENT_SF,rounded=True)
text(s,MX+Inches(0.32),y+Inches(0.25),col-Inches(0.6),Inches(0.5),[[("메신저  ",18,ACCENT,True),("✕",18,ACCENT,True)]])
text(s,MX+Inches(0.32),y+Inches(0.9),col-Inches(0.64),Inches(1.3),
     [[("상황만 전달하고 끝낸다.",14,TEXT,False)],[("\"안 된다고 합니다.\" 결정을 상사에게 떠넘긴다.",13,TEXT,False)],[("→ 담당자로서 무능해 보인다",12,MUTED,False)]],line_sp=1.2,space=6)
x2=MX+col+Inches(0.4)
rect(s,x2,y,col,Inches(2.4),fill=TEAL_SF,rounded=True)
text(s,x2+Inches(0.32),y+Inches(0.25),col-Inches(0.6),Inches(0.5),[[("플레이어  ",18,TEAL,True),("✓",18,TEAL,True)]])
text(s,x2+Inches(0.32),y+Inches(0.9),col-Inches(0.64),Inches(1.3),
     [[("상황 + 원인 + 대안을 함께 보고한다.",14,TEXT,False)],[("\"이렇게 해보겠습니다.\"",13,TEXT,False)],[("→ 채택 여부를 떠나 신뢰를 얻는다",12,MUTED,False)]],line_sp=1.2,space=6)
ry=y+Inches(2.7)
rect(s,MX,ry,CW,Inches(1.5),fill=PAPER,rounded=True)
text(s,MX+Inches(0.32),ry+Inches(0.22),CW-Inches(0.6),Inches(0.4),[("사례 · 물류 마감 사고",13,INK,True)])
text(s,MX+Inches(0.32),ry+Inches(0.68),CW-Inches(0.64),Inches(0.7),
     [[("정 사원: \"업체가 내일 발송이 어렵다고 합니다.\" → 김 팀장: \"그럼 어떻게 하지?\" → 정 사원: \"저도 잘 모르겠습니다…\"",13,TEXT,False)]],line_sp=1.25)
footer(s)

# ---- 19. M4 변환 실습 ----
ba_slide("MODULE 04 · 실습", "메신저 → 플레이어 변환",
    "메신저식 보고에 '대안 제시'를 붙여 플레이어식으로 바꾼다.",
    "[물류] \"업체가 내일은 발송이 어렵다고 합니다.\"\n\n[매출] \"매출 10% 하락. 환율 변동 탓이라 합니다.\"",
    "AFTER — 대안을 더한 보고",
    ["[물류] \"미발송분은 3일 후 출고됩니다. 긴급분만이라도",
     "  보내달라 요청하고, 안 되면 제가 직접 받아 다녀오겠습니다.\"",
     "",
     "[매출] \"…환율 영향이 적은 국내산으로 상품을",
     "  구성하는 것은 어떨까 싶습니다.\""],
    hint="대안의 채택 여부를 떠나, 대안을 들고 가는 것만으로 플레이어의 역할을 한 것이다.")

# ---- 20. M4 정리 ----
summary_slide("MODULE 04 · 정리", "보고의 자세 — 핵심",
    [[("나쁜 소식엔 ",15,RGBColor(0xE6,0xEC,0xF6),False),("항상 대안 1개 이상",15,WHITE,True),("을 함께 보고한다",15,RGBColor(0xE6,0xEC,0xF6),False)],
     [("\"잘 모르겠다\"로 끝내지 않고 ",15,RGBColor(0xE6,0xEC,0xF6),False),("내가 할 행동",15,WHITE,True),("을 말한다",15,RGBColor(0xE6,0xEC,0xF6),False)],
     [("상사가 처리하게 두지 말고 ",15,RGBColor(0xE6,0xEC,0xF6),False),("문제 해결의 주체",15,WHITE,True),("가 된다",15,RGBColor(0xE6,0xEC,0xF6),False)]])

# ---- 21. Divider 5 ----
divider("05","MODULE 05",["상사 유형 파악 &","맞춤 소통"],
        "모든 상사와 잘 지내긴 어렵다. 하지만 '어떤 스타일인지' 아는 것만으로 합을 맞추기 쉬워진다. 친밀도 × 업무관리로 4가지 유형을 나눠 본다.")

# ---- 22. M5 4유형 매트릭스 ----
s = slide()
y = header(s, "MODULE 05 · 핵심개념", "상사 유형 매트릭스 (친밀도 × 업무관리)")
col=(CW-Inches(0.4))/2
quads=[("관계중시 · 완벽주의","친밀하면서 꼼꼼함. 챙겨주되 디테일 요구.",PRIMARY_SF),
       ("관계중시 · 방목","동네 선배처럼 챙기며 믿고 맡김.",TEAL_SF),
       ("개인주의 · 완벽주의","공과 사 구분 + 성과 중심. 업무로만 신뢰.",AMBER_SF),
       ("개인주의 · 방목","간섭은 적지만 가이드도 적음. 자율과 책임.",ACCENT_SF)]
for i,(t1,d,sf) in enumerate(quads):
    r=i//2;c=i%2
    x=MX+(col+Inches(0.4))*c; ty=y+(Inches(1.9)+Inches(0.3))*r
    rect(s,x,ty,col,Inches(1.9),fill=sf,rounded=True)
    text(s,x+Inches(0.3),ty+Inches(0.28),col-Inches(0.6),Inches(0.5),[(t1,16,INK,True)])
    text(s,x+Inches(0.3),ty+Inches(0.85),col-Inches(0.6),Inches(0.9),[(d,13,TEXT,False)],line_sp=1.2)
footer(s)

# ---- 23. M5 유형별 소통법 (표 느낌) ----
s = slide()
y = header(s, "MODULE 05 · 핵심", "유형별 핵심 소통법")
rows=[("관계중시형","업무 외 상담·질문, 진심 어린 칭찬, 경조사 챙기기",PRIMARY),
      ("개인주의형","먼저 업무로 신뢰, 섣부른 사생활 공유·고민상담 자제",TEAL),
      ("완벽주의형","결론부터·잦은 중간보고, 실수 줄이고 장점 어필",AMBER),
      ("방목형","스스로 과제 프레이밍, 토론 두려워 말고 의견 제안",ACCENT)]
cy=y
for nm,d,c in rows:
    rect(s,MX,cy,CW,Inches(0.95),fill=PAPER,rounded=True)
    rect(s,MX,cy,Pt(5),Inches(0.95),fill=c)
    text(s,MX+Inches(0.35),cy+Inches(0.28),Inches(2.6),Inches(0.5),[(nm,15,c,True)])
    text(s,MX+Inches(3.2),cy+Inches(0.3),CW-Inches(3.4),Inches(0.5),[(d,13.5,TEXT,False)],anchor=MSO_ANCHOR.MIDDLE)
    cy+=Inches(1.08)
footer(s)

# ---- 24. M5 외전 주의 ----
s = slide()
y = header(s, "MODULE 05 · 외전", "소통이 어려운 3유형 — 주의")
warn=[("사조직형","관계를 이익에 악용. 라인 따라 평가가 갈린다."),
      ("마이크로매니저형","사사건건 개입·닦달. 작은 결정도 못 하게 된다."),
      ("구렁이형","앞뒤가 다르고 말 바꿈. 이야기를 이간질에 쓴다.")]
cy=y
for nm,d in warn:
    rect(s,MX,cy,CW,Inches(1.05),fill=ACCENT_SF,rounded=True)
    rect(s,MX,cy,Pt(5),Inches(1.05),fill=ACCENT)
    text(s,MX+Inches(0.35),cy+Inches(0.18),Inches(4),Inches(0.45),[(nm,16,ACCENT,True)])
    text(s,MX+Inches(0.35),cy+Inches(0.62),CW-Inches(0.7),Inches(0.4),[(d,13,TEXT,False)])
    cy+=Inches(1.2)
ry=cy+Inches(0.02)
text(s,MX,ry,CW,Inches(0.6),
     [[("→ 유형 파악 전엔 ",13,MUTED,False),("회사 사람에게 일 불만을 말하지 않기.",13,INK,True),(" 구렁이형이면 속마음 감추고 가볍게 대화, 긴밀한 관계는 멀리.",13,MUTED,False)]],line_sp=1.2)
footer(s)

# ---- 25. 종합 자가진단 ----
s = slide()
y = header(s, "WRAP-UP", "나의 보고 스킬 자가진단")
checks=["지시를 받으면 그 자리에서 목적·범위·형식을 질문한다",
        "큰 과제는 개괄조사 후 기대 수준을 상사와 먼저 맞춘다",
        "진행이 막히면 완성 전이라도 중간 보고를 한다",
        "이메일에 첨부 설명·1차 의견·추가 정보를 함께 담는다",
        "보고서는 두괄식으로, 사실에 내 의견과 근거를 더한다",
        "나쁜 소식엔 항상 대안을 1개 이상 함께 보고한다",
        "우리 상사의 유형과 선호하는 보고 방식을 안다"]
cy=y
for ck in checks:
    rect(s,MX,cy,Inches(0.32),Inches(0.32),fill=WHITE,line=PRIMARY,line_w=1.6)
    text(s,MX+Inches(0.55),cy-Inches(0.02),CW-Inches(0.6),Inches(0.4),[(ck,13.5,TEXT,False)],anchor=MSO_ANCHOR.MIDDLE)
    cy+=Inches(0.62)
footer(s)

# ---- 26. 한 장 요약 ----
s = slide()
y = header(s, "SUMMARY", "신입의 보고 5원칙")
five=[("질문","모르면 그 자리에서 묻는다. 질문은 신입의 특권.",PRIMARY),
      ("프레이밍","목적을 정의하고 기대 수준을 미리 맞춘다.",TEAL),
      ("타이밍","구두 보고는 빠르게, 자주, 진행 중에도.",AMBER),
      ("맥락+의견","이메일·보고서엔 설명과 내 의견을 담는다.",PRIMARY),
      ("플레이어","문제엔 늘 대안을 들고 간다.",ACCENT)]
cy=y
for i,(t1,d,c) in enumerate(five):
    rect(s,MX,cy,CW,Inches(0.84),fill=PAPER,rounded=True)
    rect(s,MX,cy,Inches(2.3),Inches(0.84),fill=c,rounded=True)
    text(s,MX,cy+Inches(0.2),Inches(2.3),Inches(0.5),[(t1,17,WHITE,True)],align=PP_ALIGN.CENTER)
    text(s,MX+Inches(2.6),cy+Inches(0.22),CW-Inches(2.8),Inches(0.5),[(d,14,TEXT,False)],anchor=MSO_ANCHOR.MIDDLE)
    cy+=Inches(0.96)
footer(s)

# ---- 27. 클로징 ----
s = slide(PRIMARY_DK)
rect(s,0,0,SW,SH,fill=PRIMARY_DK)
rect(s,Inches(0.9),Inches(2.7),Inches(0.9),Pt(6),fill=ACCENT)
text(s,Inches(0.9),Inches(3.0),Inches(11.5),Inches(1.5),[("결국, 보고는 신뢰입니다.",38,WHITE,True)])
text(s,Inches(0.9),Inches(4.5),Inches(11),Inches(1.0),
     [("오늘 배운 5원칙을 내일 첫 보고부터 하나씩 적용해 보세요.",17,RGBColor(0xC9,0xD6,0xEE),False)])

import os
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "보고스킬_실습_교안.pptx")
prs.save(OUT)
print(f"✅ PPTX 생성 완료: {OUT}  (슬라이드 {len(prs.slides.__iter__.__self__._sldIdLst)}장)")
